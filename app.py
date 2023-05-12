
import sys
import os
from docx import Document
from docx.shared import Pt
from docx.shared import Inches
from docx.oxml.ns import qn
import openpyxl as xl
import xlwings as xw
from tkinter import *
from PIL import Image ,ImageTk
from tkinter import messagebox
from tkinter import ttk
import random
import math
from win32com import client as wc
import win32print
import tempfile
import win32api
import pythoncom
from CeLiang import *
from tools import *
import time
from natsort import natsorted


#################################################################

def source_path(relative_path):
    if os.path.exists(relative_path):
        base_path = os.path.abspath(".")
    elif getattr(sys, 'frozen', False):
        base_path = sys._MEIPASS
    return os.path.join(base_path, relative_path)


background=source_path('background')
RES=source_path('res')
muban=source_path('muban') 
icon=source_path('icon') 


def log(*msg):
    f=open('日志.txt','a+',encoding='UTF-8')
    f.write(f'{time.strftime("%Y-%m-%d %H:%M:%S", time.localtime())}  :  {msg}\n')
    f.close()
def PJFM(a:float):
    if a<0:
        return f'左 {-a}'
    elif a==0:
        return '中'
    else:
        return f'右 {a}'

################################################################
ce=CeLiang()


class OPTION:

    def getBM(zh=0.0,filename=None):
        global res

        name=res+'\\BM' if filename ==None else filename
        data=open(name, "r+",encoding='UTF-8')
        gc={}
        for sz in data:        
            gc=eval(sz)
        near=[abs(round(i-zh,3)) for i in gc.keys()]
        key1=round(zh+min(near),3)
        key2=round(zh-min(near),3)

        if (key1) in gc:
            return (key1,gc[key1][0],gc[key1][1])
        else:
            return (key2,gc[key2][0],gc[key2][1])


    def getCtr(zh=0.0,filename=None):
        global res
        name=res+'\\ContralPoint' if filename ==None else filename
        data=open(name, "r+",encoding='UTF-8')
        KZD={}
        for v in data:        
            KZD=eval(v)
        keys=list(KZD.keys())
        near=[abs(round(i-zh,3)) for i in KZD.keys()]
        key1=round(zh+min(near),3)
        key2=round(zh-min(near),3)
        a=None
        dt=[]
        if (key1) in KZD:
            a=key1
            dt.append(KZD[key1])
        else:
            a=key2
            dt.append(KZD[key2])
        if keys.index(a)+1 <len(keys):
            dt.append(KZD[keys[keys.index(a)+1]])
        else:
            dt.append(KZD[keys[keys.index(a)-1]])
        return dt

########################################################################







class SP:

    def AngleChange(angle):
        d = int(angle)
        m = int((angle - int(angle)) * 60)
        s = round((angle - int(angle)) * 3600 - m * 60, 2)
        dms = "{0}°{1}′{2}″".format(d, m, s)
        return dms


    def mileageToStr(km):
        if km==None:
            return ''
        k = int(km // 1000)
        b = int((km - k * 1000) // 100)
        s = int((km - k * 1000 - b * 100) // 10)
        g = int((km - k * 1000 - b * 100 - s * 10) // 1)
        f = km - km // 1
        ZH = "K{0}+{1}{2}{3:0>3.03f}".format(k, b, s, g + f)
        return ZH



    def LMZD(bm=[],H=0.0,L=30.0,data=[],i=0,Hsx=None):
        
        name=bm[1]
        Hsz=bm[2]
        hd=abs(H/L*30)
        if i==0 and Hsx==None:
            if hd <4 and hd>3.5:
                sz=round(random.uniform(hd+0.7,hd+1.0) if H<0 else random.uniform(0.7,1.0),3)
                dest=round(random.uniform(hd+0.7,hd+1.0) if H>0 else random.uniform(0.7,1.0),3)
            elif hd>=4:
                hd=2.0
                sz=round(random.uniform(hd+0.7,hd+1.0) if H<0 else random.uniform(0.7,1.0),3)
                dest=round(random.uniform(hd+0.7,hd+1.0) if H>0 else random.uniform(0.7,1.0),3)         
            else:
                sz=round(random.uniform(hd+1.0,hd+1.6) if H<0 else random.uniform(1.0,1.6),3)
                dest=round(random.uniform(hd+1.0,hd+1.6) if H>0 else random.uniform(1.0,1.6),3)
                hd=2.0
            if  abs(H)<4.1 and abs(L)<30:
                Hsx=round(Hsz+sz,3)
                data.append([name,'',sz,'','',Hsx,'',Hsz,'',''])
                return (Hsx,dest,1)
            else:
                data.append([name,'',sz,'','','','',Hsz,'',''])
            Hsx=round(Hsz+sz,3)
        else:
            sz=0
            dest=0
            hd=2.0 if hd>4.0 else hd
                    
        i=1 if i==0 else i
        while round(H+sz-dest,3)!=0:
            hs=round(random.uniform(0.7+hd,1.0+hd) if H<0 else random.uniform(0.7,1.0),3)
            qs=round(random.uniform(0.7,1.0) if H<0 else random.uniform(0.7+hd,1.0+hd),3)
            if abs(round(H+sz-dest,3))<=hd and i>=int(L//30):
                qs=round(hs+H+sz-dest,3)
                if qs<0.4:
                    qst=round(random.uniform(0.7,0.7+hd),3)
                    d=round(qst-qs,3)
                    hs=round(hs+d,3)
                    qs=qst
                elif qs>4.8:
                    qst=round(random.uniform(4.5-hd,4.5),3)
                    d=round(qst-qs,3)
                    hs=round(hs+d,3)
                    qs=qst                    
                Hsx=round(Hsx+hs-qs,3)
                data.append(['ZD%d'%i,'',hs,'',qs,Hsx,'','','',''])
                break
            else:
                data.append(['ZD%d'%i,'',hs,'',qs,'','','','',''])
                Hsx=round(Hsx+hs-qs,3)
                H=round(H+hs-qs,3)
                i=i+1                           
        return (Hsx,dest,i+1)

    def ZD(bm=[],H=0.0,L=30.0,data=[],i=0,Hsx=None):
        
        name=bm[1]
        Hsz=bm[2]
        hd=abs(H/L*30) 
        if i==0 and Hsx==None:
            if hd <4 and hd>3.5:
                sz=round(random.uniform(hd+0.7,hd+1.0) if H<0 else random.uniform(0.7,1.0),3)
                dest=round(random.uniform(hd+0.7,hd+1.0) if H>0 else random.uniform(0.7,1.0),3)
            elif hd>=4:
                hd=2.0
                sz=round(random.uniform(hd+0.7,hd+1.0) if H<0 else random.uniform(0.7,1.0),3)
                dest=round(random.uniform(hd+0.7,hd+1.0) if H>0 else random.uniform(0.7,1.0),3)         
            else:
                sz=round(random.uniform(hd+1.0,hd+1.6) if H<0 else random.uniform(1.0,1.6),3)
                dest=round(random.uniform(hd+1.0,hd+1.6) if H>0 else random.uniform(1.0,1.6),3)
                hd=2.0
            if  abs(H)<4.1 and abs(L)<30:
                Hsx=round(Hsz+sz,3)
                data.append([name,sz,'','',Hsx,'',Hsz,''])
                return (Hsx,dest,1)
            else:
                data.append([name,sz,'','','','',Hsz,''])
            Hsx=round(Hsz+sz,3)
        else:
            sz=0
            dest=0
            hd=2.0 if hd>4.0 else hd
            
        i=1 if i==0 else i
        while round(H+sz-dest,3)!=0:
            hs=round(random.uniform(0.7+hd,1.0+hd) if H<0 else random.uniform(0.7,1.0),3)
            qs=round(random.uniform(0.7,1.0) if H<0 else random.uniform(0.7+hd,1.0+hd),3)
            if abs(round(H+sz-dest,3))<=hd and i>=int(L//30):
                qs=round(hs+H+sz-dest,3)                
                if qs<0.4:
                    qst=round(random.uniform(0.7,0.7+hd),3)
                    d=round(qst-qs,3)
                    hs=round(hs+d,3)
                    qs=qst
                elif qs>4.8:
                    qst=round(random.uniform(4.5-hd,4.5),3)
                    d=round(qst-qs,3)
                    hs=round(hs+d,3)
                    qs=qst
                Hsx=round(Hsx+hs-qs,3)
                data.append(['ZD%d'%i,hs,'',qs,Hsx,'','',''])
                break
            else:
                data.append(['ZD%d'%i,hs,'',qs,'','','',''])
                Hsx=round(Hsx+hs-qs,3)
                H=round(H+hs-qs,3)
                i=i+1                           
        return (Hsx,dest,i+1)

    
    def HEADER(KZD=[],*args):
        
        if KZD!=None :
            if args==():
                X0=KZD[0][1]
                Y0=KZD[0][2]
                H0=KZD[0][3]
                X2=KZD[1][1]
                Y2=KZD[1][2]
                H2=KZD[1][3]
                s=round(((X2-X0)**2+(Y2-Y0)**2)**0.5,3)
                a=math.degrees(math.atan2(X2-X0,Y2-Y0))
                α=SP.AngleChange(a) if a>0 else SP.AngleChange(a+360)
                KZD.append(s)
                KZD.append(α)
                return KZD
            else:
                X0=KZD[0][1]
                Y0=KZD[0][2]
                X1=args[0]
                Y1=args[1]
                s=round(((X1-X0)**2+(Y1-Y0)**2)**0.5,3)
                a=math.degrees(math.atan2(X1-X0,Y1-Y0))
                α=SP.AngleChange(a) if a>0 else SP.AngleChange(a+360)
                return (α,s)

        else:
            return None

    def DQSIZE(rhf=3.25,H0=5.0,mode='衡重式挡墙'):
        global res        
        h0=round(H0,2)
        file=open(res+'\\'+'DQ_size', "r+",encoding='UTF-8')
        dq={}
        for h in file:
            dq=eval(h)
        def getDQ(dq,mode,h0):
            return dq[mode][h0] if h0 in dq[mode] else None
        d=getDQ(dq,mode,h0)
        if d==None:
            return None
        #"H0":['h3','b4','b21','n','hn', 'm','m1']
        h3=d[0]
        b4=d[1]
        b21=d[2]
        n=d[3]
        hn=d[4]
        m=d[5]
        m1=d[6]
        H=round(h0-h3-hn,3)
        if mode !='路堑墙':
            ex1=round(hn*1.5,3)
            ex2=round(h3*m,3) if mode=='仰斜式挡墙' else 0
            B1=random.uniform(0.4,0.7)
            B2=random.uniform(0.9,1.4)
            jkq1=round(rhf+0.5+H*m+b21+ex2-ex1+B1,3)
            jkq2=round(jkq1-b4-B2,3)
            HQ1=-H0+B1*1.5
            HQ2=-H0+B2*1.5


            jk1=round(rhf+0.5+H*m+b21+ex2-ex1,3)
            jk2=round(jk1+ex1-b4,3)
            HH0=-H0-0.1
            jd1=round(rhf+0.5+H*m+b21+ex2,3)
            jd2=round(jd1-b4,3)
            HH1=-H0
            jcd1=round(rhf+0.5+H*m+b21,3)
            jcd2=round(jd2-(hn+h3)*m1,3)
            HH2=round(-H0+hn+h3,3)
            qsd1=round(rhf+0.5,3)
            qsd2=rhf
            HH3=0

        else:

            ex1=round(hn*1.5,3)

            jkq1=round(rhf+0.4-0.5*m1-b21+ex1,3)
            jkq2=round(rhf+0.4-0.5*m1-b21+b4,3)
            
            HQ2=round(random.uniform(-0.5,0.5),3)
            HQ1=round(-1.0-hn-h3+H0+random.uniform(-0.5,0.5),3)

            jk2=round(rhf+0.4-0.5*m1-b21+ex1,3)
            jk1=round(rhf+0.4-0.5*m1-b21+b4,3)
            HH0=round(-1.0-hn-h3,3)
            jd1=round(rhf+0.4-0.5*m1-b21,3)
            jd2=round(rhf+0.4-0.5*m1-b21+b4,3)
            HH1=round(-1.0-hn-h3,3)
            jcd1=round(rhf+0.4-0.5*m1-b21,3)
            jcd2=round(rhf+0.4-0.5*m1-b21+b4,3)
            HH2=-1
            qsd1=round(rhf+0.4-0.5*m1+H*m1,3)
            qsd2=round(rhf+0.4-0.5*m1-b21+b4,3)
            HH3=round(-1.0-hn-h3+H0,3)
        return{'基坑开挖前':([jkq1,jkq2],[HQ1,HQ2]),'基坑开挖后':([jk1,jk2],[HH0,HH0]),'基坑底':([jd1,jd2],[round(HH1+hn,3),HH1]),'基础顶':([jcd1,jcd2],[HH2,HH2]),'墙身顶':([qsd1,qsd2],[HH3,HH3])}


    def DQSIZE2(rhf=3.25,H0=5.0,mode='衡重式挡墙'):
        global res        
        h0=round(H0,1)
        file=open(res+'\\'+'DQ_size', "r+",encoding='UTF-8')
        dq={}
        for h in file:
            dq=eval(h)
        def getDQ(dq,mode,h0):
            return dq[mode][h0] if h0 in dq[mode] else None
        d=getDQ(dq,mode,h0)
        if d==None:
            return None
        #"H0":['h3','b4','b21','n','hn', 'm','m1']
        h3=d[0]
        b4=d[1]
        b21=d[2]
        n=d[3]
        hn=d[4]
        m=d[5]
        m1=d[6]
        H=round(h0-h3-hn,3)
        if mode !='路堑墙':
            ex1=round(hn*1.5,3)
            ex2=round(h3*m,3) if mode=='仰斜式挡墙' else 0

            B1=random.uniform(0.4,0.7)
            B2=random.uniform(0.9,1.4)
            jkq1=round(rhf+0.5+H*m+b21+ex2-ex1+B1,3)
            jkq2=round(jkq1+ex1-b4-B2,3)
            HQ1=-H0+B1*1.5
            HQ2=-H0+B2*1.5

            jk1=round(rhf+0.5+H*m+b21+ex2-ex1,3)
            jk2=round(jk1+ex1-b4,3)
            HH0=-H0-0.1
            jd1=round(rhf+0.5+H*m+b21+ex2,3)
            jd2=round(jd1-b4,3)
            HH1=-H0
            jcd1=round(rhf+0.5+H*m+b21,3)
            jcd2=round(jd2-(hn+h3)*m1,3)
            HH2=round(-H0+hn+h3,3)
            qsd1=round(rhf+0.5,3)
            qsd2=rhf
            HH3=0

        else:

            ex1=round(hn*1.5,3)
            jkq1=round(rhf+0.4-0.5*m1-b21+ex1,3)
            jkq2=round(rhf+0.4-0.5*m1-b21+b4,3)
            
            HQ2=round(random.uniform(-0.5,0.5),3)
            HQ1=round(-1.0-hn-h3+H0+random.uniform(-0.5,0.5),3)

            jk2=round(rhf+0.4-0.25*m1-b21+ex1,3)
            jk1=round(rhf+0.4-0.25*m1-b21+b4,3)
            HH0=round(-0.65-hn-h3,3)
            jd1=round(rhf+0.4-0.25*m1-b21,3)
            jd2=round(rhf+0.4-0.25*m1-b21+b4,3)
            HH1=round(-0.65-hn-h3,3)
            jcd1=round(rhf+0.4-0.25*m1-b21,3)
            jcd2=round(rhf+0.4-0.25*m1-b21+b4+m*(h3+hn),3)
            HH2=-0.65
            qsd1=round(rhf+0.4-0.25*m1+H*m,3)
            qsd2=round(rhf+0.4-0.25*m1-b21+b4+m*H0,3)
            HH3=round(-0.65-hn-h3+H0,3)
        return{'基坑开挖前':([jk1,jk2],[HQ1,HQ2]),'基坑开挖后':([jk1,jk2],[HH0,HH0]),'基坑底':([jd1,jd2],[round(HH1+hn,3),HH1]),'基础顶':([jcd1,jcd2],[HH2,HH2]),'墙身顶':([qsd1,qsd2],[HH3,HH3]),'回填':([jk2,round(jk2-1,3)],[0,0])}
            

    #衡重式挡墙
    #**kwargs={side:'','tp':'仰斜式挡墙','基坑底']}
    def DQGC(zh,**kwargs):
        global ce
        msg='success'
        H=abs(kwargs['size'])

        ce.SET(zh,kwargs['rw'])
        rh=abs(ce.side(kwargs['side'])[1])
        
        PJS= SP.DQSIZE(rh,H,kwargs['gongcheng']) if kwargs['project']!='会东县营盘村通村公路工程' else SP.DQSIZE2(rh,H,kwargs['gongcheng'])
        if PJS ==None:
            raise Exception('挡墙高度超出设计')
        PJ =PJS[kwargs['gongxu']] if '回填' not in  kwargs['gongxu'] else PJS['回填']
        print(PJ)
        pianju= [[zh,i,PJ[1][PJ[0].index(i)]] for i in PJ[0]] if kwargs['side']=='右侧' else [[zh,-i,PJ[1][PJ[0].index(i)]] for i in PJ[0]] 
        print(pianju)
        return pianju ,msg       


    #圆管涵
    #data=[reHigh,d,'圆管涵','工序']
    def HDGC(zh,**kwargs):
        global ce
        bd={0.75:0.11,1.0:0.12,1.5:0.14}
        δ=bd[kwargs['D']]
        k1=round(1.5*kwargs['D']+2*δ+0.3+0.5,3)
        ce.SET(zh,kwargs['rw'])

        return {'基坑开挖后':[round(zh-k1,3),round(zh+k1,3)]} ,'ok'



    #土方路基
    # zh 桩号，rw 道路宽度，LJ 路肩宽度，reHigh 相对设计路面高度。
    #data=[reHigh]
    def TFGC(zh,**kwargs):
        global ce
        ce.SET(zh,kwargs['rw'])
        wide=ce.widen
        
        ex=round(-kwargs['high']*1.5,3)
        
        if '全' in kwargs['side']:
            zhpj=[[zh,round(wide[0]-ex,3),0],[zh,0,0],[zh,round(wide[1]+ex,3),0]]
        elif kwargs['side']=='右侧':
            
            zhpj=[[zh,round(wide[1]+ex,3),0],[zh,0,0],]
        elif kwargs['side']=='左侧':
            zhpj=[[zh,round(wide[0]-ex,3),0],[zh,0,0]]

        return zhpj,'ok'

    # zh 桩号，rw 道路宽度，LJ 路肩宽度，SIDE 路肩位置。
    #data=[reHigh,SIDE]
    def LJGC(zh,**kwargs):
        global ce
        ce.SET(zh,kwargs['rw'])
        Half=ce.side(kwargs['side'])[1]
        ex=kwargs['lj'] if  kwargs['side']=='右侧' else -kwargs['lj']
        return [[zh,Half,0],[zh,round(Half-ex,3),0]] ,'ok'

    # 边沟
    
    def PSGC(zh,**kwargs):
        global ce
        ce.SET(zh,kwargs['rw'])
        Half=ce.side(kwargs['side'])[1]

        if 'Ⅰ型' in kwargs['gongcheng']:
            ex=0.65 if kwargs['side']=='右侧' else -0.65
            return [[zh,Half,0],[zh,round(Half+ex,2),0]] ,'ok'

        elif 'Ⅱ型' in kwargs['gongcheng']:
            h=random.uniform(0.4,0.8)
            ex1=round(h*1.5,3) if kwargs['side']=='右侧' else -round(h*1.5,3)
            ex2=round(round(ex1+0.9,2),3) if kwargs['side']=='右侧' else -round(round(ex1+0.9,2),3)
            return [[zh,round(Half+ex1,3),-h],[zh,round(Half+ex2,3),-h]],'ok'
        elif 'Ⅲ型' in kwargs['gongcheng']:
            ex=0.4 if kwargs['side']=='右侧' else -0.4
            return [[zh,Half,0],[zh,round(Half+ex,2),0]] ,'ok'            
        elif 'Ⅳ型' in kwargs['gongcheng']:
            ex1=-0.25 if kwargs['side']=='右侧' else 0.25
            ex2= 0.75 if kwargs['side']=='右侧' else -0.75
            return [[zh,round(Half-ex1,3),0],[zh,round(Half+ex2,3),0]] ,'ok'
    # zh 桩号，rw 道路宽度，LJ 路肩宽度
    #data=[reHigh]
    def LJJG(zh,**kwargs):
        global ce
        ce.SET(zh,kwargs['rw'])
        wide=ce.widen
        return [[zh,round(wide[0]+kwargs['lj'],3),0],[zh,0,0],[zh,round(wide[1]-kwargs['lj'],3),0]] ,'ok'



    # zh 桩号，rw 道路宽度，LJ 路肩宽度
    #data=[reHigh]
    def LMGC(zh,**kwargs):
        global ce
        LJ=kwargs['lj']
        ce.SET(zh,kwargs['rw'])
        wide=ce.widen
        return [[zh,round(wide[0]+kwargs['lj'],3),0],[zh,0,0],[zh,round(wide[1]-kwargs['lj'],3),0]] ,'ok'
        

    
    # zh 桩号，rw 道路宽度，LJ 路肩宽度    
    def Other(zh,**kwargs):
        global ce
        ce.SET(zh,kwargs['rw'])
        wide=ce.widen
        return [[zh,wide[0],0],[zh,0,0],[zh,wide[1],0]] ,'自定义'

    '''
    project:工程项目名称
    gongcheng:工程名称
    zhuanghao：桩号位置
    gongxu:工序
    data:数据[[SZ1,2.503,'','','','',2074.251,''],[ZD1,2.768,'',0.928,'','','','']]
    path:路径

    '''

    def writeFX(app,**kwargs):

        DD=kwargs['data']
        wb=app.books.open(muban+'\\全站仪放线记录表.xlsx')
        sample=wb.sheets['sheet']
        i=1
        for dt in DD:

            
            data=dt[0]
            
            num=len(data)
            yigao=round(random.uniform(1.2,1.6),3)
            pg=(num-1)//11
            header=dt[1]
            
            for p in range(1,pg+2):
                
                sample.api.Copy(Before=sample.api)
                
                sht = wb.sheets[f'sheet ({2})']
                
                new_sheet_name = f'{kwargs["zhuanghao"][:8]}数据({i}.{p})'
                
                sht.api.Name = new_sheet_name

                print('did muban\n',sht.api.Name)
                sht.range('B1').value=kwargs['project']
                sht.range('C5').value=f'{kwargs["gongcheng"]}'
                sht.range('L5').value=f'{kwargs["zhuanghao"]}{kwargs["gongxu"]}'
                sht.range('C7').value=header[0][0] #测点编号
                sht.range('F7').value=header[0][1] #测点坐标
                sht.range('F8').value=header[0][2] #测点坐标
                sht.range('F9').value=header[0][3] #测点坐标
                sht.range('I7').value=header[1][0] #后视点编号
                sht.range('M7').value=header[1][1] #后视点坐标
                sht.range('M8').value=header[1][2] #后视点坐标
                sht.range('M9').value=header[1][3] #后视点坐标
                sht.range('Q7').value=header[2]    #极坐标
                sht.range('Q9').value=header[3]    #极坐标
                sht.range('U9').value=f'{yigao}'
                
                w=(p-1)*11+11 if num>((p-1)*11+11) else num        
                sht.range('B12').value=data[(p-1)*11:w]
                
            i=i+1
            
        wb.save(kwargs['path']+f'\\{kwargs["zhuanghao"]}{kwargs["gongxu"]}-'+'放线记录.xlsx')
        wb.close()

    def writePM(app,**kwargs):

        DD=kwargs['data']
        wb=app.books.open(muban+'\\全站仪平面位置检测表.xlsx')
        sample=wb.sheets['sheet']
        i=1        
        for dt in DD:

            data=dt[0]
            num=len(data)
            yigao=round(random.uniform(1.2,1.6),3)
            pg=(num-1)//10
            header=dt[1]
            for p in range(1,pg+2):            
                
                sample.api.Copy(Before=sample.api)
                sht = wb.sheets[f'sheet ({2})']
                new_sheet_name = f'{kwargs["zhuanghao"][:8]}数据({i}.{p})'
                sht.api.Name = new_sheet_name

                sht.range('B1').value=kwargs['project']
                sht.range('C5').value=f'{kwargs["gongcheng"]}'
                sht.range('L5').value=f'{kwargs["zhuanghao"]}{kwargs["gongxu"]}' 
                sht.range('C7').value=header[0][0] #测点编号
                sht.range('F7').value=header[0][1] #测点坐标
                sht.range('F8').value=header[0][2] #测点坐标
                sht.range('F9').value=header[0][3] #测点坐标
                sht.range('I7').value=header[1][0] #后视点编号
                sht.range('L7').value=header[1][1] #后视点坐标
                sht.range('L8').value=header[1][2] #后视点坐标
                sht.range('L9').value=header[1][3] #后视点坐标
                sht.range('P7').value=header[2]    #极坐标
                sht.range('P9').value=header[3]    #极坐标
                sht.range('U9').value=yigao
                w=(p-1)*10+10 if num>((p-1)*10+10) else num        
                sht.range('B12').value=data[(p-1)*10:w]
            i=i+1
        wb.save(kwargs['path']+f'\\{kwargs["zhuanghao"]}{kwargs["gongxu"]}-'+'平面位置检测.xlsx')
        wb.close()

    def writeGC(app,**kwargs):
        DATA=kwargs['data']
        wb=app.books.open(muban+'\\高程.xlsx')
        sample=wb.sheets['sheet']
        i=1
        for dt in DATA:
            data=dt[0]
            num=len(data)
            pg=(num-1)//17
            for p in range(0,pg+1):
                sample.api.Copy(Before=sample.api)
                sht = wb.sheets[f'sheet ({2})']
                new_sheet_name = f'{kwargs["zhuanghao"][:8]}{kwargs["gongxu"][4:]}数据({i}.{p+1})'
                sht.api.Name = new_sheet_name
                sht=wb.sheets[sht.api.Name]
                sht.range('A1').value=kwargs['project']
                sht.range('B5').value=f"{kwargs['gongcheng']}"
                sht.range('E5').value=f"{kwargs['zhuanghao']}{kwargs['gongxu']}"
                w=p*17+17 if num>(p*17+17) else num
                sht.range('A9').value=data[p*17:w]                
                if p==pg:
                    sht.range('B26').value=f"ΣH=H测-H设={dt[1][0]}-{dt[1][1]}={dt[1][2]}mm，符合精度要求"
            i=i+1
        wb.save(kwargs['path']+f"\\{kwargs['zhuanghao']}{kwargs['gongxu']}-"+'高程检测.xlsx')
        wb.close()


    def writeGCJS(app,**kwargs):
        DATA=kwargs['data']
        wb=app.books.open(muban+'\\路基路面高程检测、计算.xlsx')
        sample=wb.sheets['sheet']
        i=1
        for dt in DATA:
            data=dt[0]
            num=len(data)
            pg=(num-1)//19
            for p in range(0,pg+1):
                sample.api.Copy(Before=sample.api)
                sht = wb.sheets[f'sheet ({2})']
                new_sheet_name = f'{kwargs["zhuanghao"][:8]}{kwargs["gongxu"][4:]}数据({i}.{p+1})'
                sht.api.Name = new_sheet_name
                sht=wb.sheets[sht.api.Name]
                sht.range('A1').value=kwargs['project']
                sht.range('B6').value=f"{kwargs['zhuanghao']}{kwargs['gongcheng']}"
                sht.range('F6').value=f"{kwargs['gongxu']}"
                w=p*19+19 if num>(p*19+19) else num
                sht.range('A10').value=data[p*19:w]                
            i=i+1
        wb.save(kwargs['path']+f"\\{kwargs['zhuanghao']}{kwargs['gongxu']}-"+'路基路面高程检测、计算.xlsx')
        wb.close()

############################################################################            

##############################################################
       
class APP:
    global Page

    global resources
    global folders
    Page=None

    ENGS={'土方工程':'TFGC','涵洞工程':'HDGC','挡墙工程':'DQGC','路肩工程':'LJGC','排水工程':'PSGC','路基工程':'LJJG','路面工程':'LMGC','自定义':'Other'}
    resources={(fd if  os.path.isdir(source_path('res')+'//'+fd) else None):fd for fd in os.listdir(source_path('res'))}
    GC_list={}

    
    #########初始化###################################
    def __init__(self):
        global root
        global title
        global image
        global pic
        global bground
        global bg_name
        global Page
        global project
        global engineering
        global resources

        root=Tk()
        root.title('小工具 Version stable 1.3')
        root.geometry('1280x760')
        root.attributes("-alpha", 0.98)
        root.iconbitmap(icon+'\\'+'favicon.ico')


        mbar=Menu(root)
        menu_main=Menu(mbar,tearoff=0)
        menu_main.add_command(label='主页',command=self.MainPage)
        menu_main.add_command(label='批量替换',command=self.view_rep_SP)
        menu_main.add_command(label='文件夹内批量替换',command=self.view_Normal)
        menu_main.add_command(label='文档转换',command=self.view_exchange)
        menu_main.add_command(label='批量打印',command=self.Printer)
        menu_main.add_command(label='坐标查询',command=self.lookup)

        mbar.add_cascade(label='主菜单',menu=menu_main)
        
        Pro=Menu(mbar,tearoff=0)
        project=StringVar()
        for xm in resources:
            Pro.add_radiobutton(label=xm ,variable=project,value=xm,command=self.SETTING)

        mbar.add_cascade(label='项目名称',menu=Pro)

        Class=Menu(mbar,tearoff=0)

        engineering=StringVar()
        for gc in self.ENGS:
            Class.add_radiobutton(label=gc,variable=engineering ,value=gc,command=self.Deal)

        mbar.add_cascade(label='类别',menu=Class)


        setting=Menu(mbar,tearoff=0)
        setting.add_command(label='换肤',command=self.skin)
        
        mbar.add_cascade(label='设置',menu=setting)
        mbar.add_command(label='关于')
        mbar.add_command(label='VIP',command=self.VIP)

        mbar.add_command(label='退出',command=root.quit)
        root.config(menu=mbar)
        bg_name=StringVar()
        bg_name.set(background+'\\bg.jpg')
        image=Image.open(f'{bg_name.get()}') if os.path.isfile(f'{bg_name.get()}') else None
        pic=ImageTk.PhotoImage(image) if os.path.isfile(f'{bg_name.get()}') else None
        msg=''
        bground=Label(root,text=msg,justify=LEFT,compound = CENTER,image=pic)
        bground.pack(side=LEFT)
       
        Page=PanedWindow(root,orient=VERTICAL,width=1000)
        Page.place(x=50,y=60)

        self.Home()
        root.bind("<Configure>",self.background_image)
        

        root.mainloop()
    ###########背景调整################################
    def background_image(self,event):
        global root
        global image
        global pic
        global bground
        if os.path.isfile(f'{bg_name.get()}') :
            image=Image.open(f'{bg_name.get()}').resize((int(root.winfo_width()),int(root.winfo_height())))
            pic=ImageTk.PhotoImage(image)
            bground['image']=pic
    ############背景切换############################
    def image_change(self):
        global root
        global image
        global pic
        global bground
        image=Image.open(background+f'\\{bg_name.get()}').resize((int(root.winfo_width()),int(root.winfo_height())))
        pic=ImageTk.PhotoImage(image)
        bground['image']=pic
    ########################################
    def skin(self):
        global image
        global pic
        global bground
        bgs=[]
        for filename in os.listdir(background):
            bgs.append(filename)

        bg_name.set(f'{bgs[random.randint(0,len(bgs)-1)]}')
        image=Image.open(background+f'\\{bg_name.get()}').resize((int(root.winfo_width()),int(root.winfo_height())))
        pic=ImageTk.PhotoImage(image)
        bground['image']=pic
        ##################################################
    def SETTING(self):

        global project
        global res
        global ce

        if project.get() in resources:
            res=RES+'\\'+resources[project.get()]
            with open(RES+"\\"+resources[project.get()]+'\\projectlist.txt',"r",encoding="utf-8") as f:
                GC_list=eval(f.read())
                print(GC_list)
            INIT=open(res+'\\'+'init.txt',"r+", encoding="UTF-8")
            data=[]
            for i in INIT:
                data=eval(i)
            log(data)
            log(res)
            ce=CeLiang(res,data[0],data[1],data[2],data[3],data[4],data[5],data[6],data[7]) if data!=[] else None
            log(ce.res)
        else:
            ce=None
        self.Home()

    def Deal(self):
        global engineering
        obj=SP
        self.fun=getattr(obj,self.ENGS[engineering.get()])
        if 'tree' in globals():
            del globals()['tree']

        self.MainPage()

    ########清除###########################################
    def Clear(self,Page:PanedWindow):
        
        if Page!=None:
            for  i in Page.panes():
                Page.remove(i)
                i=None
    ########主页#########################################
    def Home(self):
        global Page
        global project
        self.Clear(Page)

        Page.add(Label(text='欢迎',font=('仿宋',20)))
        Page.add(Label(text='注：高程，平面位置需先设置，相关内容',font=('仿宋',13,'bold'),fg='red'))
        if project.get()=='会东县人民村通村公路工程':
            Page.add(Label(text='人民村路：0~6100路肩为0.25；6100~9645.5路肩为0.4，9645.5~10581.786为0.5；\n0~9645.5路面宽度为6.5，9645.5~10581.786为路面宽度7.5。',font=('仿宋',16),fg='brown'))        
        if 'tree' in globals():
            del globals()['tree']
    def MainPage(self):
        global Page
        global project
        global engineering

        self.Clear(Page)
        if project.get()!='' and engineering.get()!='':
            Page.add(Button(text='放线记录',command=self.FX_view,height=5,width=40))
            Page.add(Button(text='平面位置检测',command=self.PMWZ_view,height=5,width=40))
            Page.add(Button(text='高程检测',command=self.HeightCheck,height=5,width=40))
        if 'tree' in globals():
            del globals()['tree']

    def lookup(self):
        global Page
        global project
        
        if project.get() =='':
            return self.Home()
        self.Clear(Page)
        Page.add(Label(text='查询坐标高程',font=('仿宋 25')))
        pan1=PanedWindow()
        pan1.add(Label(text='桩号：',font=('仿宋 20')))
        zh=DoubleVar()
        pan1.add(Entry(textvariable=zh))
        Page.add(pan1)
        pan2=PanedWindow()
        pan2.add(Label(text='偏距：',font=('仿宋 20'),fg='red'))
        pj=DoubleVar()
        pan2.add(Entry(textvariable=pj))
        
        Page.add(pan2)
        rw=DoubleVar()
        pan3=PanedWindow()
        pan3.add(Label(text='路宽：',font=('仿宋 20'),fg='red'))
        pan3.add(Entry(textvariable=rw))
        Page.add(pan3)
        Lb=Label(text='',font=('楷体 18'),height=9)
        Page.add(Lb)


        def GP(lb,zh0,pj0,rw):
            global ce
            ce.SET(zh0,rw)
            P=ce.Point(pj0)
            H=ce.Height(pj0)[1]
            lb['text']=f'高程：{H} \n坐标：{P}'

        Page.add(Button(text='查询',command=lambda:GP(Lb,zh.get(),pj.get(),rw.get())))


    ########桩号替换#########################################
    def view_rep_SP(self):
        global root
        global Page        
        global values
        global HOME
        self.Clear(Page)
        
        if Page not in globals():
            Page=PanedWindow(root,orient=VERTICAL)
            Page.place(x=60,y=60)
        Page.add(Label(text='特殊（桩号）批量替换工具',font=('仿宋 22')))
        pan1=PanedWindow()
        pan1.add(Label(text='需要处理的文件路径：'))
        path=StringVar()
        path.set('mode')
        pan1.add(Entry(textvariable=path))
        Page.add(pan1)
        pan2=PanedWindow()
        pan2.add(Label(text='需要处理的桩号文件：'))
        data=StringVar()
        data.set('rep.xlsx')
        pan2.add(Entry(textvariable=data))
        Page.add(pan2)
        values=[]
        PATH=None

        def deal():
            global PATH
            global SHEET
            try:
                if  os.path.isdir(path.get()):
                    PATH=path.get()  
                else:
                    raise Exception('未指定路径')
                if  os.path.isfile(data.get()):
                    dt=xl.load_workbook(data.get())
                    SHEET=list(dt.worksheets[0].rows)
                else:
                    raise Exception('未找到文件')
                row=list(dt.worksheets[0].rows)[0]
                if values==[]:
                    for cell in row:
                        pan=PanedWindow()
                        var1=StringVar()
                        pan.add(Label(text=f'替换{cell.column}（旧）：'))
                        var1.set(cell.value)
                        pan.add(Entry(textvariable=var1))
                        values.append(var1)
                        
                        Page.add(pan)
                    Page.add(Button(text='开始替换',command=self.replace_SP))
                dt.close()
            except Exception as err:

                messagebox.showerror('showerror', err)

        Page.add(Button(text='提交',command=deal))
            
    def replace_SP(self):
        global values
        global PATH
        global SHEET
        try:
            app=xw.App(visible=False, add_book=False)
            for row in SHEET:
                folder=' '+row[0].value
                
                if os.path.isdir(folder):
                    pass
                else:
                    os.makedirs(folder)

                for filename in natsorted(os.listdir(PATH)):
                    if os.path.splitext(filename)[1] in ['.docx']:
                        doc1=Document(f'{PATH}\\{filename}')
                        for cell in row:
                            VALUE=cell.value if cell.value!=None else ' '
                            Tool.replace_word(values[cell.column-1].get(),VALUE,doc1,)
                        doc1.save(f'{folder}\\{row[0].value}-{filename}')
                        
                        
                    elif os.path.splitext(filename)[1] in ['.xlsx','xls']:
                        wb=app.books.open(f'{PATH}\\{filename}')
                        for cell in row:
                            VALUE=cell.value if cell.value!=None else ' '
                            Tool.replace_excel(values[cell.column-1].get(),VALUE,wb)
                        wb.save(f'{folder}\\{row[0].value}-{filename}')
                        wb.close()
         
        except Exception as err:
            messagebox.showerror('错误', err)

        else:
            messagebox.showinfo('信息', '成功')
        finally:
            app.quit()

    #########普通批量 替换###########################################
    def view_Normal(self):
        global Page
        global PATH
        global values
        self.Clear(Page)
        Page.add(Label(text='指定路径批量替换',font=('仿宋 22')))
        opt1=PanedWindow()
        opt1.add(Label(text='需要处理的文件夹路径：'))

        PATH=StringVar()
        opt1.add(Entry(textvariable=PATH))
        
        Page.add(opt1)
        values=[]

        try:

            def add():
                opt2=PanedWindow()
                opt2.add(Label(text='旧'))
                old=StringVar()
                opt2.add(Entry(textvariable=old))
                opt2.add(Label(text='新'))
                new=StringVar()
                opt2.add(Entry(textvariable=new))
                values.append([old,new])
                INDEX=values.index([old,new])
                def close(pan,index):
                    Page.remove(pan)
                    values.pop(index)
                opt2.add(Button(text='×',command=lambda:close(opt2,INDEX)))
                Page.add(opt2)
        except Exception as e:
            messagebox.showerror('错误', e)


        opt3=PanedWindow()
        
        opt3.add(Button(text='增加',command=add))
        Page.add(opt3)

        Page.add(Button(text='开始替换',command=self.replace_N))
        pass
    def replace_N(self):
        global values
        global PATH

        try:
            app=xw.App(visible=False, add_book=False)
            for filename in natsorted(os.listdir(PATH.get())):

                if os.path.splitext(filename)[1] in ['.docx']:
                    doc1=Document(f'{PATH.get()}\\{filename}')
                    for row in values:
                        Tool.replace_word(row[0].get(),row[1].get(),doc1)

                    doc1.save(f'{PATH.get()}\\{filename}')
                    
                elif os.path.splitext(filename)[1] in ['.xlsx','.xls']:
                    wb=app.books.open(f'{PATH.get()}\\{filename}')
                    for row in values:
                        Tool.replace_excel(row[0].get(),row[1].get(),wb)
                    wb.save(f'{PATH.get()}\\{filename}')
                    wb.close()
        except Exception as e:
            messagebox.showerror('错误', e)
        else:
            messagebox.showinfo('信息', '成功')
        finally:
            app.quit()
            
    ##########格式转换#############################################
    def view_exchange(self):
        global Page
        global PATH
        global PATH2        
        self.Clear(Page)
        Page.add(Label(text='Word、Excel 转换成 .docx,.xlsx',font=('仿宋 22')))
        suffix=['doc','docx','txt','xml','html','dot','pdf','xps','csv','xls','xlsx','dif','dbf']

        fmt={'doc':0,'dot':1,'txt':2,'换行 txt':3,'dos txt':4,'dos 换行 txt':5,'rtf':6,'Unicode txt':7,\
        'html':8,'单个 html':9,'Filtered html':10,'xml':11,'document xml':12,'document macro xml':13,\
        'template xml':14,'template macro xml':15,'docx':16,'pdf':17,'xps':18,\
        'excel pdf':0,'csv':6,'dbase 2 dbf':7,'dbase 3 dbf':8,'dif':9,'dbase 4 dbf':11,'excel html':44,\
        'xlt':17,'Macintosh txt':19,'Windows txt':20,'dos txt':21,'Macintosh csv':22,\
        'Windows csv':23,'dos csv':24,'xlsx':51,'97-2003 xls':56,'utf8 csv':62,\
        }
        args=[]
        PATH=StringVar()
        PATH2=StringVar()
        def getFMT(event):
            print (fmt[cob.get()])
            arg1=[]
            arg2=[]           
            for suff in suffix :
                if suff in cob.get():
                    arg1=[suff,fmt[cob.get()]]
                if suff in cob2.get():
                    arg2=[suff,fmt[cob2.get()]]
            print(arg1)
            print(arg2)
            self.args=tuple(arg1+arg2)
            

        opt1=PanedWindow()
        opt1.add(Label(text='需要处理的文件夹路径：'))
        
        opt1.add(Entry(textvariable=PATH),minsize=350)


            
        opt1.add(Label(text='word格式'))
        code=StringVar()
        cob=ttk.Combobox(values=list(fmt.keys())[:19],textvariable=code)
        cob.bind("<<ComboboxSelected>>",getFMT)
        opt1.add(cob,minsize=350)
        Page.add(opt1)

        opt2=PanedWindow()
        opt2.add(Label(text='excel格式'))
        code2=StringVar()
        cob2=ttk.Combobox(values=list(fmt.keys())[19:],textvariable=code2)
        cob2.bind("<<ComboboxSelected>>",getFMT)
        opt2.add(cob2,minsize=350)

        opt2.add(Label(text='存储路径：'))
        
        opt2.add(Entry(textvariable=PATH2),minsize=350) 
        Page.add(opt2)
        Page.add(Button(text='开始转换',command=lambda:self.exchange(PATH,PATH2,*self.args)))
       
        
    def exchange(self,PATH,PATH2,*args):
        if os.path.isdir(PATH.get()):
            if os.path.isdir(PATH2.get()):
                path=PATH2.get()
            else:
                os.makedirs(PATH2.get())
                path=PATH2.get()
            print(PATH.get())
            Tool.zhuanhuan(PATH.get(),PATH2.get(),*args)
    #############批量打印#############################################
    def Printer(self):
        global Page
        global opt_md
        global show_md
        global show_data
        self.Clear(Page)
        show_md=False
        Page.add(Label(text='批量打印',font=('仿宋',18,'bold'),fg='purple'))
        opt1=PanedWindow()
        opt1.add(Label(text='目标文件夹'))
        PATH=StringVar()
        opt1.add(Entry(textvariable=PATH))
        opt_list=None
        opt1.add(Button(text='获取文件',command=lambda:self.getFiles(opt_list,PATH)))
        opt1.add(Button(text='获取文件夹',command=lambda:self.getFolders(opt_list,PATH)))
        Page.add(opt1)
        opt_list=PanedWindow(height=300,width=500)
        Page.add(opt_list)

        opt2=PanedWindow()
        opt2.add(Label(text='打印份数：'))
        num=IntVar()
        num.set(1)
        opt2.add(Entry(textvariable=num))
        opt2.add(Label(text='每份结束等待（秒）：'))
        sl=IntVar()
        sl.set(3)
        opt2.add(Entry(textvariable=sl))
        opt2.add(Button(text='打印',command=lambda:self.PRINT(num,sl)))
        Page.add(opt2)
        opt_md=PanedWindow(height=30)
        Page.add(opt_md)

    def getFiles(self,top,path):
        global tree
        global mode

        if 'tree' not in globals():
            mode='file'
            tree=ttk.Treeview(top,show="headings")
            s=ttk.Style()
            s.theme_use('default')
            tree['columns']=['路径','文件名']
            tree.column("路径",width=100,anchor="center")
            tree.column("文件名",width=100,anchor="center")
            tree.heading("路径",text="路径")
            tree.heading("文件名",text="文件名")
            i=1
            files=os.listdir(path.get())        
            for filename in natsorted(files):
                if os.path.isfile(path.get()+'\\'+filename):
                    tree.insert('',i,values=(path.get(),filename))
                    i=i+1
            tree.bind("<Delete>",self.Del)                 
            top.add(tree)
        else:
            i=1
            files=os.listdir(path.get())        
            for filename in natsorted(files):
                if os.path.isfile(path.get()+'\\'+filename):
                    tree.item() if tree.exists(i) else tree.insert('',i,values=(path.get(),filename))
                    i=i+1
            for item in tree.get_children()[i+1:]:
                if tree.exists(item):
                    tree.delete(item)                        
    def getFolders(self,top,path):
        global tree
        global mode
        if 'tree' not in globals():
            mode='folder'
            tree=ttk.Treeview(top,show="headings")
            s=ttk.Style()
            s.theme_use('default')
            tree['columns']=['路径','子文件夹名']
            tree.column("路径",width=100,anchor="center")
            tree.column("子文件夹名",width=100,anchor="center")
            tree.heading("路径",text="路径")
            tree.heading("子文件夹名",text="子文件夹名")
            i=0
            files=os.listdir(path.get())

            for filename in natsorted(files):
                if os.path.isdir(path.get()+'\\'+filename):
                    k=len(natsorted(os.listdir(f'{path.get()}\\{filename}')))
                    tree.insert('','end',f'ID{i}',values=(path.get(),filename))
                    j=1
                    for file in natsorted(os.listdir(f'{path.get()}\\{filename}')):
                        tree.insert(f'ID{i}','end',f'print{i}.{j}',values=(f'{path.get()}\\{filename}',file))
                        j=j+1
                    i=i+1

            tree.bind("<Delete>",self.Del)
            
            top.add(tree)
            
        else:
            i=0
            files=os.listdir(path.get())

            for filename in natsorted(files):
                if os.path.isdir(path.get()+'\\'+filename):
                    k=len(natsorted(os.listdir(f'{path.get()}\\{filename}')))
                    tree.item(f'ID{i}' ,values=(path.get(),filename)) if  tree.exists(f'ID{i}') else tree.insert('','end',f'ID{i}',values=(path.get(),filename))
                    j=1
                    for file in natsorted(os.listdir(f'{path.get()}\\{filename}')):
                        tree.item(f'print{i}.{j}' ,values=(f'{path.get()}\\{filename}',file)) if tree.exists(f'print{i}.{j}') else  tree.insert(f'ID{i}','end',f'print{i}.{j}',values=(f'{path.get()}\\{filename}',file))
                        j=j+1
                    for item in tree.get_children()[j+1:]:
                        if tree.exists(item):
                            tree.delete(item)                    
                    i=i+1
            for item in tree.get_children()[i+1:]:
                if tree.exists(item):
                    tree.delete(item)                                
    def PRINT(self,num,sl):
        global tree
        global mode

        try:

            SL=int(sl.get()) if int(sl.get())>3 else 3
            NUM=int(num.get())

            if mode=='folder':
                for item in tree.get_children():
                    for n in range(NUM):
                        log(f'开始打印第{n+1}份')
                        for file in tree.get_children(item):
                            filename='\\'.join(tree.item(file,'values'))
                            Tool.Printer(filename)
                            time.sleep(SL/2)
                        time.sleep(SL)
            elif mode=='file':
                for n in range(NUM):
                    log(f'开始打印第{n+1}份')
                    for file in tree.get_children():
                        filename='\\'.join(tree.item(file,'values'))
                        
                        Tool.Printer(filename)
                        time.sleep(SL)
        except Exception as e:
            messagebox.showerror('错误',f'打印失败：{e}')
        else:
            messagebox.showinfo('信息','打印结束')

    ###########平面位置###################################################
    def View(self,title,JC,tb_hd,font1,rowHigh,cmds):
        global project
        global Page
        global pulldown
        global engineering
        global Pane3
        global side
        global show_md
        global Pane_mod
        global tree
        global JianCe
        JianCe=JC

        show_md=False
        Page.add(Label(text=title,font=('仿宋',22,'bold'),bg='darkblue',fg='white',height=2),pady=3)
        gongxu_list={'挡墙工程':['基坑开挖前','基坑开挖后','基坑底','基础顶','墙身顶','墙背回填第   层'],'土方路基':['上路床','下路床','上路提','下路堤'],\
        '涵洞工程':['基坑开挖前','管形基础顶面','端墙顶面','跌水井井口','总体','流水面','基坑底','基础顶面','台身顶面','盖板顶面','翼墙顶面','八字墙顶面','跌水井井底','跌水井井口','台背填土第   层'],\
        '路基工程':['级配碎石垫层底面','级配碎石垫层顶面','水泥稳定碎石底基层底面','水泥稳定碎石底基层顶面','水泥稳定碎石基层底面','水泥稳定碎石基层顶面'],\
        '排水工程':['基坑开挖前','基坑开挖后','基坑底','基础顶','沟身顶面'],\
        '路肩工程':['路肩底面','路肩顶面'],'路面工程':['沥青混凝土下面层底面','沥青混凝土下面层顶面','沥青混凝土上面层底面','沥青混凝土上面层顶面'],\
        '自定义':['自定义涵洞','基坑开挖前','管形基础顶面','端墙顶面','跌水井井口','总体','流水面','基坑底','基础顶面','台身顶面','盖板顶面','翼墙顶面','八字墙顶面','跌水井井底','跌水井井口','台背填土第   层']\
        }

        properties={'project':['项目名称',list(resources.keys()),project],\
        'gongcheng':['工程名称：',self.GC_list[engineering.get()],StringVar()],\
        'zh1':['工程部位',None,DoubleVar()],\
        'zh2':['--',None,DoubleVar()],\
        'side':['位置',['左侧','右侧'],StringVar()],\
        'gongxu':['工序',gongxu_list[engineering.get()],StringVar()],\
        'rw':['路面宽度',None,DoubleVar()],
        'lj':['路肩宽度',None,DoubleVar()],\
        'ping':['频率',None,DoubleVar()],\
        'high':['相对高度',None,DoubleVar()],\
        'size':['手输墙高',None,DoubleVar()],\
        'pc_s':['偏差小（mm）',None,IntVar()],\
        'pc_b':['偏差大（mm）',None,IntVar()],\
        } 
        if engineering.get() not in ['挡墙工程','排水工程','路肩工程']:
            properties.pop('side')
        if engineering.get()in ['涵洞工程','自定义']:
            properties['size'][0]='手输管径'
            properties.pop('zh2')
        kw={k:properties[k][2] for k in properties} 
        kw['lj'].set(0.25)
        kw['rw'].set(6.5)
        
        
        i=0
        pan=None
        for dd in properties.values():
            if i%2==0:
                pan=PanedWindow(Page)
            LB= Label(text=dd[0],bg='gray',fg='white',font=font1)   
            pan.add(LB,minsize=100,padx=3,pady=3)
            con=Entry(textvariable=dd[2],font=font1) if dd[1]==None else ttk.Combobox(values=dd[1],textvariable=dd[2],font=font1,state='readonly')
            pan.add(con,minsize=350,padx=3,pady=3)
            if i%2==0:
                Page.add(pan)
            i=i+1

        Pane_data=None
        Pane5=PanedWindow()
        Pane5.add(Button(text='自动获取',relief=GROOVE,command=lambda:self.Tree(Pane_data,tb_hd,**kw)),minsize=200,padx=10,pady=5)
        Pane5.add(Button(text='修改计算',relief=GROOVE,command=lambda:self.calculate(tree,**kw)),minsize=200,padx=10,pady=5)
        Page.add(Pane5)
        Page.add(Label(text='++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++'))
        Pane_mod=PanedWindow(height=rowHigh)
        Page.add(Pane_mod)
        Pane_data=PanedWindow(height=250)
        Page.add(Pane_data)
        Pane_save=PanedWindow()
        Pane_save.add(Label(text='保存位置：'))
        PATH=StringVar()
        Pane_save.add(Entry(textvariable=PATH,width=30))
        Pane_save.add(Button(text='保存(SAVE)',relief=GROOVE,command=lambda:cmds(tree,PATH,**kw,)))
        Page.add(Pane_save)
        

    
    def FX_view(self):
        global Page
        
        self.Clear(Page)

        self.View('放线记录','FXJL',['桩号','偏距','高程','X','Y','计算方位角','计算距离'],('楷体',12),25,self.fxDeal)
    
    def fxDeal(self,tree,PATH,**kwargs):
        try:
            app=xw.App(visible=False, add_book=False)
            kwargs={k:kwargs[k].get() for k in kwargs}
            ZH1=kwargs['zh1']

            path=f"{PATH.get()}\\{kwargs['gongcheng']}\\{SP.mileageToStr(ZH1)}{ '-'+SP.mileageToStr(kwargs['zh2']) if 'zh2' in kwargs else ''}{kwargs['side'] if 'side' in kwargs else ''}{kwargs['gongxu']}"

            if os.path.isdir(path):
                pass
            else:
                os.makedirs(path)
            log(path)
            kwargs['path']=path
            DATA=[]
            HEAD=None
            data=[]
            No=1
            kwargs['zhuanghao']=f"{SP.mileageToStr(ZH1)}{ '-'+SP.mileageToStr(kwargs['zh2']) if 'zh2' in kwargs else ''} {kwargs['side'] if 'side' in kwargs else ''}"

            ZHS=tree.get_children()
            for item in ZHS:
                row =tree.item(item,'values')

                zh=round(float(row[0]),3)
                pianju=row[1]
                X=round(float(row[3]),4)
                Y=round(float(row[4]),4)

                a,s=SP.HEADER(OPTION.getCtr(zh),*(X,Y))
                pc=round(random.uniform(-0.02,0.02),3)
                s1=round(s+pc,3)
                H=round(float(row[2]),3)
                
                if HEAD!=SP.HEADER(OPTION.getCtr(zh)):

                    if data!=[]:
                        log('data',data)
                        DATA.append([data,HEAD])
                        data=[]
                    
                    data.append([f'{SP.mileageToStr(zh)} {pianju}',None,X,None,Y,None,a,None,None,s,None,s1,None,abs(pc),None,H,None])
                else:
                    
                    data.append([f'{SP.mileageToStr(zh)} {pianju}',None,X,None,Y,None,a,None,None,s,None,s1,None,abs(pc),None,H,None])
                HEAD=SP.HEADER(OPTION.getCtr(zh))

            if data!=[]:
                log('data',data)
                DATA.append([data,HEAD])
            kwargs['data']=DATA
            SP.writeFX(app,**kwargs)
        except Exception as e:
            messagebox.showerror('错误信息',f'发生错误,{e}，执行失败！！')
        else:
            messagebox.showinfo('信息', '成功')        
        finally:
            app.quit()
        
    ###########高程####################################################
    def HeightCheck(self):
        global Page
        
        self.Clear(Page)

        self.View('高程检测','GCJC',['桩号','偏距','设计高程'],('楷体',12),25,self.gaochengDeal)


    def gaochengDeal(self,tree,PATH,**kwargs):

        try:
            app=xw.App(visible=False, add_book=False)
            kwargs={k:kwargs[k].get() for k in kwargs}
            ZH1=kwargs['zh1']
            path=f"{PATH.get()}\\{kwargs['gongcheng']}\\{SP.mileageToStr(ZH1)}{ '-'+SP.mileageToStr(kwargs['zh2']) if 'zh2' in kwargs else ''}{kwargs['side'] if 'side' in kwargs else ''}{kwargs['gongxu']}"
            if os.path.isdir(path):
                pass
            else:
                os.makedirs(path)
            kwargs['path']=path
            kwargs['zhuanghao']=f"{SP.mileageToStr(ZH1)}{ '-'+SP.mileageToStr(kwargs['zh2']) if 'zh2' in kwargs else ''} {kwargs['side'] if 'side' in kwargs else ''}" if '回填' not in kwargs['gongxu'] else f"{SP.mileageToStr(ZH1)}{ '-'+SP.mileageToStr(kwargs['zh2']) if 'zh2' in kwargs else ''}"
            DATA=[]
            data=[]
            BM=None
            Hsx=None
            dest=None
            Hq=None
            zhq=None
            i=0
            ZHS=tree.get_children()
            for item in ZHS:
                row=tree.item(item,'values')

                zh=round(float(row[0]),3)
                pianju=row[1]
                HS=round(float(row[2]),3)
                pc=random.randint(kwargs['pc_s'],kwargs['pc_b'])
                if BM!=OPTION.getBM(zh):
                    if data !=[]:
                        
                        
                        H=Hsx-BM[2]-dest
                        L=BM[0]-zhq
                        back=SP.ZD(BM,H,abs(L),data,i,Hsx) if abs(L)!=0 else SP.ZD(BM,H,abs(L+2),data,i,Hsx)
                        Hsx=back[0]
                        Ep=random.randint(-1*i,1*i)
                        data.append([f'{BM[1]}','','',round(Hsx-BM[2],3),'',round(BM[2]+Ep/1000,3),BM[2],Ep])                        
                        
                        DATA.append([data,[round(BM[2]+Ep/1000,3),BM[2],Ep]])
                        SP.writeGC(app,**kwargs)
                        data=[]               
                    BM=OPTION.getBM(zh)                    
                    H=BM[2]-HS
                    L=abs(zh-BM[0])
                    Hsx=None
                    i=0
                    back=SP.ZD(BM,H,abs(L),data,i,Hsx)
                    Hsx=back[0]
                    i=back[2]
                    dest=back[1]
                    zhq=zh
                    Hq=HS 
                    data.append([f'{SP.mileageToStr(zh)},{pianju}','','',round(Hsx-round(HS+pc/1000,3),3),'',round(HS+pc/1000,3),HS,pc])
                else:
                    HC=Hsx-HS            
                    L=zh-zhq
                    if HC<0.6 or HC>4.8:
                        H=Hsx-HS-dest if HC> 4.8 else HC-2
                        back=SP.ZD(BM,H,abs(L),data,i,Hsx) if abs(L)!=0 else SP.ZD(BM,H,abs(L+2),data,i,Hsx)
                        log(back)
                        Hsx=back[0]
                        i=back[2]
                        zhq=zh
                        Hq=HS        
                    data.append([f'{SP.mileageToStr(zh)},{pianju}','','',round(Hsx-round(HS+pc/1000,3),3),'',round(HS+pc/1000,3),HS,pc])
            
            log(Hq,BM,zhq,i,dest)
            
            
            H=Hsx-BM[2]-dest
            L=BM[0]-zhq
            back=SP.ZD(BM,H,abs(L),data,i,Hsx) if abs(L)!=0 else SP.ZD(BM,H,abs(L+2),data,i,Hsx)
            Hsx=back[0]
            if data!=[]:
                
                Ep=random.randint(-1*i,1*i)
                data.append([f'{BM[1]}','','',round(Hsx-BM[2],3),'',round(BM[2]+Ep/1000,3),BM[2],Ep])         
                DATA.append([data,[round(BM[2]+Ep/1000,3),BM[2],Ep]])
            kwargs['data']=DATA
            SP.writeGC(app,**kwargs)
            
        except Exception as err:
            messagebox.showerror('错误信息',f'发生错误,{err}，执行失败！！')
        else:

            messagebox.showinfo('信息', '成功')
        finally:
            app.quit()
            

    def PMWZ_view(self):
        global Page
        self.Clear(Page)
        

        self.View('平面位置检测','PMWZ',['桩号','偏距','X','Y'],('楷体',12),25,self.pmwzDeal)

    def pmwzDeal(self,tree,PATH,**kwargs):

        try:
            app=xw.App(visible=False, add_book=False)
            kwargs={k:kwargs[k].get() for k in kwargs}
            ZH1=kwargs['zh1']
            path=f"{PATH.get()}\\{kwargs['gongcheng']}\\{SP.mileageToStr(ZH1)}{ '-'+SP.mileageToStr(kwargs['zh2']) if 'zh2' in kwargs else ''}{kwargs['side'] if 'side' in kwargs else ''}{kwargs['gongxu']}"
            
            if os.path.isdir(path):
                pass
            else:
                os.makedirs(path)
            log(path)
            HEAD=None
            DATA=[]
            data=[]
            kwargs['path']=path
            kwargs['zhuanghao']=f"{SP.mileageToStr(ZH1)}{ '-'+SP.mileageToStr(kwargs['zh2']) if 'zh2' in kwargs else ''} {kwargs['side'] if 'side' in kwargs else ''}" if '回填' not in kwargs['gongxu'] else f"{SP.mileageToStr(ZH1)}{ '-'+SP.mileageToStr(kwargs['zh2']) if 'zh2' in kwargs else ''}"

            ZHS=tree.get_children()
            for item in ZHS:
                row =tree.item(item,'values')

                zh=round(float(row[0]),3)
                pianju=row[1]
                X=round(float(row[2]),4)
                Y=round(float(row[3]),4)
                pc=int(round(((25**2)/2)**0.5))
                px=random.randint(-pc,pc)
                py=random.randint(-pc,pc)
                ps=round((px**2+py**2)**0.5)
                args=()
                if HEAD!=SP.HEADER(OPTION.getCtr(zh),*args):

                    if data!=[]:
                        log('data',data)
                        DATA.append([data,HEAD])

                        SP.writePM(app,**kwargs)
                        data=[]
                    
                    data.append([f'{SP.mileageToStr(zh)} {pianju}',None,X,None,Y,None,round(X+px/1000,4),None,round(Y+py/1000,4),None,None,px,None,py,None,ps,None])
                else:
                    
                    data.append([f'{SP.mileageToStr(zh)} {pianju}',None,X,None,Y,None,round(X+px/1000,4),None,round(Y+py/1000,4),None,None,px,None,py,None,ps,None])
                HEAD=SP.HEADER(OPTION.getCtr(zh),*args)

            if data!=[]:
                log('data',data)
                DATA.append([data,HEAD])
            kwargs['data']=DATA
            SP.writePM(app,**kwargs)

       
        except Exception as e:
            messagebox.showerror('错误信息',f'发生错误,{e}，执行失败！！')
        else:
            messagebox.showinfo('信息', '成功')        
        finally:

            app.quit()
            pass
        


    ####################################################################
    def Tree(self,top,header,**kwargs):
        global tree
        global JianCe
        global ce

        kwargs={k:kwargs[k].get() for k in kwargs}
        log('其他',kwargs)
        log('方法',self.fun)
        zh1=kwargs['zh1']
        zh2=kwargs['zh2'] if 'zh2' in kwargs else None
        ping=kwargs['ping']
        kwargs.pop('zh1')
        kwargs.pop('zh2')  if 'zh2' in kwargs else None
        kwargs.pop('ping')
        msg=None
        if kwargs['gongcheng'] in ['钢筋混凝土圆管涵','钢筋混凝土圆暗板涵']:
            CDS=[round(zh1+ping*i+ping*random.uniform(0.3,0.8),3) for i in range(6)]
            pass
        else:

            num =1 if ping==0 else int((zh2-zh1)//ping)
            CDS=[round(zh1+ping*i+ping*random.uniform(0.3,0.8),3) for i in range(num)]
        
        if 'tree' in globals():
            NO=0
            index=0
            for zh in CDS:
                dd,msg=self.fun(zh,**kwargs)
                log(dd)
                if dd!=None:
                    k=len(dd)
                    i=0
                    for d in dd:
                        ce.SET(d[0],kwargs['rw'])
                        if kwargs['gongcheng'] in self.GC_list['挡墙工程']:
                            
                            b=ce.side(kwargs['side'])[1]
                            height=ce.Height(d[1])[1]+kwargs['high']+d[2]
                            point=ce.Point(d[1])
                        else:
                            height=ce.Height(d[1])[1]+kwargs['high']
                            point=ce.Point(d[1])
                        if d[1]==0:
                            pj='中'
                        else:
                            pj=f'左,{-d[1]}' if d[1]<0 else f'右,{d[1]}'
                        if JianCe=='PMWZ':
                            values=[d[0],pj,f'{point[0]:.4f}',f'{point[1]:.4f}']
                        elif JianCe=='GCJC':
                            values=[d[0],pj,f'{height:.3f}']
                        elif JianCe =='FXJL':
                            KZD=OPTION.getCtr(d[0])
                            a,s=SP.HEADER(KZD,*point)
                            values=[d[0],pj,f'{height:.3f}',f'{point[0]:.4f}',f'{point[1]:.4f}',a,s]
                        NO=index*k+i
                        tree.item(f'{NO}',values=values) if tree.exists(f'{index*k+i}') else tree.insert('','end',f'{index*k+i}',values=values)
                        i=i+1
                    index=index+1
            for item in tree.get_children()[NO+1:]:
                if tree.exists(item):
                    tree.delete(item)

        else:
            tree=ttk.Treeview(show='headings')#show='headings'

            style=ttk.Style()
            style.theme_use('default')
            
            tree["columns"] = header
            for head in header:
                tree.column(f"{head}", width=80,anchor="center")
                tree.heading(f"{head}", text=f"{head}")
            
            index=0
            for zh in CDS:
                dd,msg=self.fun(zh,**kwargs)
                log(dd)                
                if dd!=None:
                    k=len(dd)
                    i=0
                    for d in dd:
                        ce.SET(d[0],kwargs['rw'])
                        if kwargs['gongcheng'] in self.GC_list['挡墙工程']:
                            b=ce.side(kwargs['side'])[1]
                            height=ce.Height(d[1])[1]+kwargs['high']+d[2]
                            point=ce.Point(d[1])
                        else:                        
                            height=ce.Height(d[1])[1]+kwargs['high']
                            point=ce.Point(d[1])
                        if d[1]==0:
                            pj='中'
                        else:
                            pj=f'左,{-d[1]}' if d[1]<0 else f'右,{d[1]}'
                        if JianCe=='PMWZ':
                            values=[d[0],pj,f'{point[0]:.4f}',f'{point[1]:.4f}']
                        elif JianCe=='GCJC':
                            values=[d[0],pj,f'{height:.3f}']
                        elif JianCe =='FXJL':
                            KZD=OPTION.getCtr(d[0])
                            a,s=SP.HEADER(KZD,*point)
                            values=[d[0],pj,f'{height:.3f}',f'{point[0]:.4f}',f'{point[1]:.4f}',a,s]
                        
                        tree.insert('','end',f'{index*k+i}',values=values)
                        
                        i=i+1
                    index=index+1            
            tree.bind("<Delete>",self.Del)
            tree.bind("<Double-1>",self.edit) 
            top.add(tree)
            messagebox.showinfo('信息',f'{msg}')
        

    ############################################################
    def Del(self,event):
        global tree
        
        for item in tree.selection():
            if tree.exists(item):
                tree.delete(item)


    def edit(self,event):

        global tree
        global mods
        global Pane_mod
        global show_md
        for item in tree.selection():
            #item = I001
            log('id:',item)
            item_text = tree.item(item, "values")

            def save(Item):
                global show_md
                values=[s.get() for s in mods]
                tree.item(Item, text="", values=values)
                tree.update()
                self.Clear(Pane_mod)
                show_md=False
                log('id:',Item)
                
            if show_md==False:
                mods=[]
                if len(item_text)>3:
                    item_text=item_text[:3]
                    
                    lbs=list(tree['columns'])
                    lbs=lbs[:3]

                else:
                    lbs=tree['columns']
                i=0
                for var in item_text:
                    tem=StringVar()
                    tem.set(var)
                    entryedit = Entry(textvariable=tem,width=10)
                    Pane_mod.add(Label(text=f'{lbs[i]}：'))
                    Pane_mod.add(entryedit)
                    mods.append(tem)
                    i=i+1
                Pane_mod.add(Button(text='保存',command=lambda:save(item)))
                show_md=True
            

    def calculate(self,tree,**kwargs):
        global JianCe
        
        global ce
        try:
            kwargs={k:kwargs[k].get() for k in kwargs}
            if kwargs['rw']!='' and kwargs['high']!='':
                rw=kwargs['rw']
                h0=kwargs['high']
                for item in tree.get_children():
                    values=tree.item(item,'values')
                    zh=round(float(values[0]),3)
                    ce.SET(zh,rw)
                    if '右' in values[1].split(','):
                        pj=round(float(values[1].split(',')[1]),3)
                    elif '左' in values[1].split(','):
                        pj=round(-float(values[1].split(',')[1]),3)
                    else:
                        pj=0
                    H=round(ce.Height(pj)[1]+h0,3)
                    point=ce.Point(pj)

                    data=['' for i in range(len(tree['columns']))]
                    data[0]=values[0]
                    data[1]=values[1]
                    

                    if JianCe=='GCJC':
                        data[2]=H
                    elif JianCe=='PMWZ':
                        data[2]=f'{point[0]:.4f}'
                        data[3]=f'{point[1]:.4f}'
                    
                    elif JianCe =='FXJL':
                        KZD=OPTION.getCtr(zh)
                        a,s=SP.HEADER(KZD,*point)
                        data[2]=values[2]
                        data[3]=f'{point[0]:.4f}'
                        data[4]=f'{point[1]:.4f}'
                        data[5]=f'{a}'
                        data[6]=f'{s}'
                        

                    tree.item(item, text="", values=data)
                    tree.update()
            else:
                raise Exception('请检查相关参数是否输入完整')
        
        except Exception as e:
            messagebox.showerror('错误', f'计算失败,发生错误：{e}')
        

    def VIP(self):
        global Page
        global project
        self.Clear(Page)
        Page.add(Label(text='批量生成数据',font=('仿宋 22')))
        ll=['基坑开挖前','基坑开挖后','基坑底','基础顶','墙身顶','墙背回填']
        properties={'file':['文件：',None,StringVar()],\
        'project':['项目名称',None,project],\
        'gongcheng':['工程名称：',list(self.ENGS.keys()),StringVar()],\
        'gongxu':['工序',ll,StringVar()],\
        'rw':['路面宽度',None,DoubleVar()],\
        'lj':['路肩宽度',None,DoubleVar()],\
        'ping':['频率',None,DoubleVar()],\
        'high':['相对高度',None,DoubleVar()],\
        'pc_s':['偏差小（mm）',None,IntVar()],\
        'pc_b':['偏差大（mm）',None,IntVar()],\
        'pw':['偏位（mm）',None,IntVar()],\
        'jgc':['结构层厚度（m）',None,DoubleVar()],\
        'save':['保存位置：',None,StringVar()],\
        } 
        kw={k:properties[k][2] for k in properties} 
        kw['rw'].set(6.5)
        kw['lj'].set(0.25)
        desk=os.path.expandvars ('C:\\Users\\%USERNAME%\\Desktop')
        kw['file'].set(f'{desk}\\test.xlsx')
        kw['save'].set(desk)
        i=0
        pan=None
        for dd in properties.values():
            if i%2==0:
                pan=PanedWindow(Page)
            LB= Label(text=dd[0],bg='gray',fg='white',font='仿宋')   
            pan.add(LB,minsize=100,padx=3,pady=3)
            con=Entry(textvariable=dd[2],font='仿宋') if dd[1]==None else ttk.Combobox(values=dd[1],textvariable=dd[2],font='仿宋')
            pan.add(con,minsize=350,padx=3,pady=3)
            if i%2==0:
                Page.add(pan)
            i=i+1
        shw=None
        btn=None
        Page.add(Button(text='获取',command=lambda:self.getlist(shw,btn,**kw)))
        shw=PanedWindow(height=300,width=500)
        Page.add(shw)
        btn=Button(text='生成数据',command=lambda:self.DEAL(btn,**kw))   
        
        Page.add(btn)

    def getlist(self,top,bn,**kwargs):
        global tree

        kwargs={k:kwargs[k].get() for k in kwargs}
        app=xw.App(visible=False, add_book=False)
        wb=app.books.open(kwargs['file'])
        sht=wb.sheets[0]
        cell = sht.used_range.last_cell
        rows = cell.row
        columns = cell.column
        heads=sht.range((1,1),(1,columns)).value
        log(heads)



        if 'tree' not in globals():            
            tree=ttk.Treeview(top,show="headings")
            s=ttk.Style()
            s.theme_use('default')
            tree['columns']=heads
            for head in heads:
                tree.column(head,width=100,anchor="center")
                tree.heading(head,text=head)
            dts=sht.range((2,1),(rows,columns)).value
            j=1
            for dt in dts:
                tree.insert('','end',f'{kwargs["gongcheng"]}{j}',values=dt)
                j=j+1
            tree.bind("<Delete>",self.Del)
            top.add(tree)
        app.quit()

    def DEAL(self,bn,**kwargs):
        if '填' not in kwargs['gongxu'].get():
            self.dealBat(**kwargs)
        else:
            if '土方' not in kwargs['gongcheng'].get():
                self.dealHT(**kwargs)
            else:
                self.dealTFHT(**kwargs)

    def dealBat(self,**kwargs):
        global tree
        global ce
        
        kws={k:kwargs[k].get() for k in kwargs}
        methd=self.ENGS[kws['gongcheng']]

        F=getattr(SP,methd)
        
        kw={}
        kw['project']=kws['project']
        kw['gongxu']=kws['gongxu']
        kw['rw'] = kws['rw']
        kw['lj'] = kws['lj']
        kw['path']=kws['save']
        kw['pc_s']=kws['pc_s']
        kw['pc_b']=kws['pc_b']
        kw['pw']=kws['pw']
        DD=[]
        name=None
        if kw['gongxu'] =='基坑开挖前':
            Funs=[self.FXbat]
        else:
            Funs=[self.PMWZbat,self.GCbat]
        for Fun in Funs:
            N=0
            for item in tree.get_children():
                try:
                    vals=tree.item(item,'values')

                    if DD!=[] and name!=vals[0]:

                        Fun(DD,**kw)
                        DD=[]
                    kw['buwei']=vals[0]
                    name=vals[0]
                    zh1=round(float(vals[1]),3)
                    zh2=round(float(vals[2]),3)
                    if kws['ping']==0.0:
                        N=3
                        if round(float(vals[6]),3)>30:
                            a=(round(float(vals[6]),3)-30)//20 
                            N= 3+a+1 if (round(float(vals[6]),3)-30)//10%2!=0 else 3+a
                        n=int(round((zh2-zh1)/round(float(vals[6]),3)*N,0))
                        if n==0:
                            n=1
                            N=N-1
                        else:
                            n=n

                        stp=(zh2-zh1)/n
                    else:
                        if (zh2-zh1)>abs(kws['ping']):
                            n=int((zh2-zh1)//abs(kws['ping']))
                        else:
                            raise Exception('频率超出范围')
                        stp=abs(kws['ping'])
                    log(stp)
                    if Fun == self.FXbat and kws['ping']<0.0:
                        zhs=[zh1,zh2]
                        
                    else:
                        zhs=[round(zh1+stp*(i-1)+random.uniform(0.3,0.8)*stp,3) for i in range(1,n+1)]

                    log(zhs)
                    

                    kw['side']=vals[3]
                    kw['gongcheng']=vals[4]
                    kw['size']=round(float(vals[7]),1)
                    log(f'kw:{kw}')
                    for zh in zhs:
                        dd,msg = F(zh,**kw)
                        ce.SET(zh,kw['rw'])
                        b=ce.side(kw['side'])[1]
                        h=ce.Height(b)[1]+kws['high']
                        dd=[[d[0],d[1],round(h+d[2],3)] for d in dd]
                        DD.extend(dd)
                        log(dd)
                except Exception as err:

                    log(err)
                else:
                    f=open('信息.txt','a+',encoding='UTF-8')
                    f.write(f'{time.strftime("%Y-%m-%d %H:%M:%S", time.localtime())}  :  {vals}成功\n')
                    f.close()
                finally:
                    pass
                if DD !=[]:
                    
                    Fun(DD,**kw)
                    
    def dealHT(self,**kwargs):

        global tree
        global ce

        kws={k:kwargs[k].get() for k in kwargs}

        methd=self.ENGS[kws['gongcheng']]


        F=getattr(SP,methd)
        
        kw={}
        kw['project']=kws['project']
        kw['gongxu']=kws['gongxu']
        kw['rw'] = kws['rw']
        kw['lj'] = kws['lj']
        kw['path']=kws['save']
        kw['pc_s']=kws['pc_s']
        kw['pc_b']=kws['pc_b']
        kw['jgc']=kws['jgc']
        kw['pw']=kws['pw']
        DD=[]
        HH=[]
        name=None
        Fun=self.HTbat

        for item in tree.get_children():
            
            vals=tree.item(item,'values')
            if HH!=[] and name!=vals[0]:
                    
                    HH.sort()
                    print(HH)
                    if '路堑墙'==vals[4]:
                        raise Exception('路堑墙无回填')                    
                    HT=[]
                    h=0                    
                    for itm in HH:
                        H=round(itm[0]-h,3)
                        i=H//0.2
                        i=i
                        HT.extend([0.2 for j in range(int(i))])
                        if round(H-0.2*i,3)>=0.1:                       
                            HT.append(round(H-0.2*i,3)) 
                        else:
                            HT.pop()
                            HT.extend([round(H-0.2*i+0.2,3)/2,round(H-0.2*i+0.2,3)/2])                        
                        h=itm[0]
                    print(HT)

                    H0=0
                    NO=len(HT)
                    for ht in HT:
                        try:
                            qu=[]
                            for hh in HH:
                                if hh[0]>=H0:
                                    qu.append([hh[1],hh[2]])

                            print('qu',qu)
                            qu.sort()
                            print(f'sort:{H0} 第{NO}层 ',qu)
                            zh1=min(qu)[0][0]
                            zh2=max(qu)[0][1]
                            hl=round(zh2-zh1,3)
                            dm=5 if hl<=30 else (hl-30)//10+5
                            stp=round(hl/dm,3)
                            zhs=[round(zh1+random.uniform(0.3,0.7)*stp+stp*(i-1),3) for i in range(1,int(dm)+1)]
                            DD=[]
                            for zh in zhs:
                                
                                h0=0
                                for hh in HH:
                                    if zh>hh[1][0] and zh<=hh[1][1]:
                                        kw['size']=hh[2]
                                        h0=hh[0]-H0
                                dd,msg=F(zh,**kw)
                                ce.SET(zh,kw['rw'])
                                bs=ce.side(kw['side'])[1]                            
                                b1=dd[0][1]-h0*0.25 if kw['gongcheng']!='衡重式挡墙' or H0>kw['size']/3 else bs-H0*0.5
                                b2=0 if H0<0.5 else dd[0][1]-h0*2/3
                                print('dd:',dd)
                                b=round(random.uniform(b2,b1),3)

                                hc=round(ce.Height(b)[1]+kws['high']-kw['jgc']-H0,3)
                                DD.append([zh,b,hc])
                            log(f'DD: No.{NO}',DD)
                            kw['NO']=NO
                            Fun(DD,**kw)
                        except Exception as er:
                            log(f"{kw['buwei']}{NO}失败：{er}")

                        else:
                            f=open('信息.txt','a+',encoding='UTF-8')
                            f.write(f"{time.strftime('%Y-%m-%d %H:%M:%S', time.localtime())}  :  {kw['buwei']} 第{NO}层 回填  成功\n")
                            f.close()
                        finally:
                            H0=round(H0+ht,3)
                            NO=NO-1


                    
                    DD=[]
                    HH=[]
                    
            kw['buwei']=vals[0]
            name=vals[0]
            zh1=round(float(vals[1]),3)
            zh2=round(float(vals[2]),3)
            l=round(float(vals[5]),3)
            L=round(float(vals[6]),3)
            kw['side']=vals[3]
            kw['gongcheng']=vals[4]
            kw['size']=round(float(vals[7]),1)            
            ht=round(float(vals[7])-kw['jgc'],3) if kw['gongcheng'] in ['仰斜式挡墙','护肩墙'] else round(float(vals[7])-kw['jgc'],3)    
            HH.append([ht,[zh1,zh2],float(vals[7])])
                    

        if HH!=[]:
            HH.sort()
            print(HH)
            if '路堑墙'==vals[4]:
                raise Exception('路堑墙无回填')                    
            HT=[]
            h=0                    
            for itm in HH:
                H=round(itm[0]-h,3)
                i=H//0.2
                i=i
                HT.extend([0.2 for j in range(int(i))])
                if round(H-0.2*i,3)>=0.1:                       
                    HT.append(round(H-0.2*i,3)) 
                else:
                    HT.pop()
                    HT.extend([round(H-0.2*i+0.2,3)/2,round(H-0.2*i+0.2,3)/2])                        
                h=itm[0]
            print(HT)

            H0=0
            NO=len(HT)
            for ht in HT:
                try:
                    qu=[]
                    for hh in HH:
                        if hh[0]>=H0:
                            qu.append([hh[1],hh[2]])

                    print('qu',qu)
                    qu.sort()
                    print(f'sort:{H0} 第{NO}层 ',qu)
                    zh1=min(qu)[0][0]
                    zh2=max(qu)[0][1]
                    hl=round(zh2-zh1,3)
                    dm=5 if hl<=30 else (hl-30)//10+5
                    stp=round(hl/dm,3)
                    zhs=[round(zh1+random.uniform(0.3,0.7)*stp+stp*(i-1),3) for i in range(1,int(dm)+1)]
                    DD=[]
                    hf=round(random.uniform(0.5,0.8),3)
                    for zh in zhs:
                        h0=0
                        for hh in HH:
                            if zh>hh[1][0] and zh<=hh[1][1]:
                                kw['size']=hh[2]
                                h0=hh[0]-H0
                        dd,msg=F(zh,**kw)
                        ce.SET(zh,kw['rw'])
                        bs=ce.side(kw['side'])[1]                            
                        b1=dd[0][1]-h0*0.25 if kw['gongcheng']!='衡重式挡墙' or H0>kw['size']/3 else bs-H0*0.5
                        b2=0 if H0<hf else dd[0][1]-h0*2/3
                        print('dd:',dd)
                        b=round(random.uniform(b2,b1),3)

                        hc=round(ce.Height(b)[1]+kws['high']-kw['jgc']-H0,3)
                        DD.append([zh,b,hc])
                    log(f'DD: No.{NO}',DD)
                    kw['NO']=NO
                    Fun(DD,**kw)
                except Exception as er:
                    log(f"{kw['buwei']}{NO}失败：{er}")

                else:
                    f=open('信息.txt','a+',encoding='UTF-8')
                    f.write(f"{time.strftime('%Y-%m-%d %H:%M:%S', time.localtime())}  :  {kw['buwei']} 第{NO}层 回填  成功\n")
                    f.close()
                finally:
                    H0=round(H0+ht,3)
                    NO=NO-1
                     
        pass



    def HTFC(self,H,h,zh1,stp,dm,fq,F,Fun,**kw):
        n=h//0.3-1 if round(h-(h//0.3)*0.3,4)>0 and round(h-(h//0.3)*0.3,4)<0.1 else h//0.3
        print(n,h)
        cs=[0.3 for i in range(int(n))]
        if round(h-(h//0.3)*0.3,4)>0 and round(h-(h//0.3)*0.3,4)<0.1:
            cs.extend([round((h+0.3-(h//0.3)*0.3)/2,3),round((h+0.3-(h//0.3)*0.3)/2,3)])  
        elif round(h-(h//0.3)*0.3,4)!=0: 
            cs.append(round(h-(h//0.3)*0.3,3))
        kw['gongxu']=fq+kw['gongxu']
        hi=0
        No=1
        kw['total']=len(cs)
        print(cs)
        for h0 in cs[::-1]:
            DD=[]
            zhs=[round(zh1+random.uniform(0.3,0.7)*stp+stp*(i-1),3) for i in range(1,int(dm)+1)]
            
            for zh in zhs:
                kw['high']=H-hi
                dd,msg=F(zh,**kw)
                ce.SET(zh,kw['rw'])
                
                d=[[i[0],i[1],round(ce.Height(i[1])[1]+kw['high']-kw['jgc'],3)] for i in dd]
                DD.extend(d)
            kw['NO']=No

            
            Fun(DD,**kw)
            H=H-hi
            No=No+1

    def HTCL(self,h,zh1,stp,dm,fq,F,Fun,**kw):

        kw['gongxu']=fq+'顶'

        DD=[]
        zhs=[round(zh1+random.uniform(0.3,0.7)*stp+stp*(i-1),3) for i in range(1,int(dm)+1)]
        
        for zh in zhs:
            kw['high']=h
            dd,msg=F(zh,**kw)
            ce.SET(zh,kw['rw'])
            
            d=[[i[0],i[1],round(ce.Height(i[1])[1]+kw['high']-kw['jgc'],3)] for i in dd]
            DD.extend(d)
        
        Fun(DD,**kw)
        
        


    def dealTFHT(self,**kwargs):

        global tree
        global ce

        kws={k:kwargs[k].get() for k in kwargs}
        methd=self.ENGS[kws['gongcheng']]
        F=getattr(SP,methd)
        
        kw={}
        kw['project']=kws['project']
        kw['gongxu']=kws['gongxu']
        kw['rw'] = kws['rw']
        kw['lj'] = kws['lj']
        kw['path']=kws['save']
        kw['pc_s']=kws['pc_s']
        kw['pc_b']=kws['pc_b']
        kw['jgc']=kws['jgc']
        kw['pw']=kws['pw']
        Type={0.3:"上路床",0.8:"下路床",1.5:"上路提"}
        DD=[]
        HH=[]
        name=None
        Fun=self.LMbat
        print('路基土石方')
        for item in tree.get_children():
            vals=tree.item(item,'values')

            stp=kws['ping']
            l=round(float(vals[5]),3)
            dm=2 if l<100 else l//stp
            zh1=round(float(vals[1]),3)
            h=round(float(vals[7]),3)
            kw['side']=vals[3]
            kw['gongcheng']=vals[4]
            kw['buwei']=vals[0]
            dd=[]

            if h>=0.1 and h<=0.3:

                self.HTFC(h,h,zh1,stp,dm,'上路床',F,Fun,**kw)
                self.HTCL(0,zh1,stp,dm,'上路床',F,self.PMWZbat,**kw)
                self.HTCL(0,zh1,stp,dm,'上路床',F,self.GCbat,**kw)


            elif h<=0.8:
                self.HTFC(0.3,0.3,zh1,stp,dm,'上路床',F,Fun,**kw)
                self.HTFC(h,round(h-0.3,3),zh1,stp,dm,'下路床',F,Fun,**kw)

                self.HTCL(0,zh1,stp,dm,'上路床',F,self.PMWZbat,**kw)
                self.HTCL(0.3,zh1,stp,dm,'下路床',F,self.PMWZbat,**kw)
                
                self.HTCL(0,zh1,stp,dm,'上路床',F,self.GCbat,**kw)
                self.HTCL(0.3,zh1,stp,dm,'下路床',F,self.GCbat,**kw)
                
                pass

            elif h<=1.5:
                self.HTFC(0.3,0.3,zh1,stp,dm,'上路床',F,Fun,**kw)
                self.HTFC(0.8,round(0.8-0.3,3),zh1,stp,dm,'下路床',F,Fun,**kw)                
                self.HTFC(h,round(h-0.8,3),zh1,stp,dm,'上路提',F,Fun,**kw)

                self.HTCL(0,zh1,stp,dm,'上路床',F,self.PMWZbat,**kw)
                self.HTCL(0.3,zh1,stp,dm,'下路床',F,self.PMWZbat,**kw)                
                self.HTCL(0.8,zh1,stp,dm,'上路提',F,self.PMWZbat,**kw)
                
                self.HTCL(0,zh1,stp,dm,'上路床',F,self.GCbat,**kw)
                self.HTCL(0.3,zh1,stp,dm,'下路床',F,self.GCbat,**kw)                
                self.HTCL(0.8,zh1,stp,dm,'上路提',F,self.GCbat,**kw)
                
                pass


            else:
                

                self.HTFC(0.3,0.3,zh1,stp,dm,'上路床',F,Fun,**kw)
                self.HTFC(0.8,round(0.8-0.3,3),zh1,stp,dm,'下路床',F,Fun,**kw)                
                self.HTFC(1.5,round(1.5-0.8,3),zh1,stp,dm,'上路提',F,Fun,**kw)
                self.HTFC(h,round(h-1.5,3),zh1,stp,dm,'下路提',F,Fun,**kw)


                self.HTCL(0,zh1,stp,dm,'上路床',F,self.PMWZbat,**kw)
                self.HTCL(0.3,zh1,stp,dm,'下路床',F,self.PMWZbat,**kw)                
                self.HTCL(0.8,zh1,stp,dm,'上路提',F,self.PMWZbat,**kw)
                self.HTCL(1.5,zh1,stp,dm,'下路提',F,self.PMWZbat,**kw)
                
                self.HTCL(0,zh1,stp,dm,'上路床',F,self.GCbat,**kw)
                self.HTCL(0.3,zh1,stp,dm,'下路床',F,self.GCbat,**kw)                
                self.HTCL(0.8,zh1,stp,dm,'上路提',F,self.GCbat,**kw)
                self.HTCL(1.5,zh1,stp,dm,'下路提',F,self.GCbat,**kw)
                
                pass



    def GCbat(self,DD,**kwargs):
        global ce
        try:

            app=xw.App(visible=False, add_book=False)

            
            path=f"{kwargs['path']}\\{kwargs['gongcheng']}\\{kwargs['gongxu']}"
            if os.path.isdir(path):
                pass
            else:
                os.makedirs(path)
            kwargs['path']=path
            condtions=['回填','路床','路提','基层','底基层','沥青']
            kwargs['zhuanghao']=f"{kwargs['buwei']} {kwargs['side'] if 'side' in kwargs else ''}" if all(con not in kwargs['gongxu'] for con in condtions) else f"{kwargs['buwei']} "
            DATA=[]
            data=[]
            BM=None
            Hsx=None
            dest=None
            Hq=None
            zhq=None
            No=1
            i=0        
            for row in DD:
                zh=round(float(row[0]),3)
                pianju=PJFM(row[1])
                HS=round(float(row[2]),3)
                pc=random.randint(kwargs['pc_s'],kwargs['pc_b'])
                if BM!=OPTION.getBM(zh):
                    if data !=[]:
                        H=Hsx-BM[2]-dest
                        L=BM[0]-zhq
                        back=SP.ZD(BM,H,abs(L),data,i,Hsx) if abs(L)!=0 else SP.ZD(BM,H,abs(L+2),data,i,Hsx)
                        Hsx=back[0]
                        Ep=random.randint(-1*i,1*i)
                        data.append([f'{BM[1]}','','',round(Hsx-BM[2],3),'',round(BM[2]+Ep/1000,3),BM[2],Ep])
                        
                        DATA.append([data,[round(BM[2]+Ep/1000,3),BM[2],Ep]])
                        data=[]
                    BM=OPTION.getBM(zh)                    
                    H=BM[2]-HS
                    L=abs(zh-BM[0])
                    Hsx=None
                    i=0
                    back=SP.ZD(BM,H,abs(L),data,i,Hsx)
                    Hsx=back[0]
                    i=back[2]
                    dest=back[1]
                    zhq=zh
                    Hq=HS 
                    data.append([f'{SP.mileageToStr(zh)},{pianju}','','',round(Hsx-round(HS+pc/1000,3),3),'',round(HS+pc/1000,3),HS,pc])
                else:
                    HC=Hsx-HS            
                    L=zh-zhq

                    if HC<0.6 or HC>4.8 or L>30:
                        H=Hsx-HS-dest if HC> 4.8 else HC-2
                        back=SP.ZD(BM,H,abs(L),data,i,Hsx) if abs(L)!=0 else SP.ZD(BM,H,abs(L+2),data,i,Hsx)
                        log(back)
                        Hsx=back[0]
                        i=back[2]
                        zhq=zh
                        Hq=HS        
                    data.append([f'{SP.mileageToStr(zh)},\n{pianju}','','',round(Hsx-round(HS+pc/1000,3),3),'',round(HS+pc/1000,3),HS,pc])
            
            log(Hq,BM,zhq,i,dest)
            #dest=random.uniform(0.6,1.5) if Hq-BM[2]>0 else random.uniform(2.6,4.5)
            
            H=Hsx-BM[2]-dest
            L=BM[0]-zhq
            back=SP.ZD(BM,H,abs(L),data,i,Hsx) if abs(L)!=0 else SP.ZD(BM,H,abs(L+2),data,i,Hsx)
            Hsx=back[0]
            if data!=[]:
                Ep=random.randint(-1*i,1*i)
                data.append([f'{BM[1]}','','',round(Hsx-BM[2],3),'',round(BM[2]+Ep/1000,3),BM[2],Ep])            
                DATA.append([data,[round(BM[2]+Ep/1000,3),BM[2],Ep]])
            kwargs['data']=DATA
            SP.writeGC(app,**kwargs)
        except Exception as err:
            log(err)
        else:
            f=open('高程信息.txt','a+',encoding='UTF-8')
            f.write(f"{time.strftime('%Y-%m-%d %H:%M:%S', time.localtime())}  :  {kwargs['zhuanghao']} 高程 成功\n")
            f.close()            
        finally:
            app.quit()
        pass

    def HTbat(self,DD,**kwargs):
        global ce
        try:

            app=xw.App(visible=False, add_book=False)            
            path=f"{kwargs['path']}\\{kwargs['gongcheng']}\\{kwargs['gongxu']}"
            if os.path.isdir(path):
                pass
            else:
                os.makedirs(path)
            kwargs['path']=path
            condtions=['回填','路床','路提','基层','底基层','沥青']
            kwargs['zhuanghao']=f"{kwargs['buwei']} {kwargs['side'] if 'side' in kwargs else ''}" if all(con not in kwargs['gongxu'] for con in condtions) else f"{kwargs['buwei']} "
            kwargs['gongxu']=f"{kwargs['gongxu']}第{kwargs['NO']}层"
            DATA=[]
            data=[]
            BM=None
            Hsx=None
            dest=None
            Hq=None
            zhq=None
            No=1
            i=0        
            for row in DD:
                
                zh=round(float(row[0]),3)
                pianju=PJFM(row[1])
                HS=round(float(row[2]),3)
                pc=random.randint(kwargs['pc_s'],kwargs['pc_b'])
                if BM!=OPTION.getBM(zh):
                    if data !=[]:
                        H=Hsx-BM[2]-dest
                        L=BM[0]-zhq
                        back=SP.ZD(BM,H,abs(L),data,i,Hsx) if abs(L)!=0 else SP.ZD(BM,H,abs(L+2),data,i,Hsx)
                        Hsx=back[0]
                        Ep=random.randint(-1*i,1*i)
                        data.append([f'{BM[1]}','','',round(Hsx-BM[2],3),'',round(BM[2]+Ep/1000,3),BM[2],Ep])
                        
                        DATA.append([data,[round(BM[2]+Ep/1000,3),BM[2],Ep]])
                        data=[]
                    BM=OPTION.getBM(zh) 
                    H=BM[2]-HS
                    L=abs(zh-BM[0])
                    Hsx=None
                    i=0
                    back=SP.ZD(BM,H,abs(L),data,i,Hsx)
                    Hsx=back[0]
                    i=back[2]
                    dest=back[1]
                    zhq=zh
                    Hq=HS 
                    data.append([f'{SP.mileageToStr(zh)}\n{pianju}','','',round(Hsx-round(HS+pc/1000,3),3),'',round(HS+pc/1000,3),HS,pc])
                else:
                    HC=Hsx-HS            
                    L=zh-zhq
                    if HC<0.6 or HC>4.8 or L>30:
                        H=Hsx-HS-dest if HC> 4.8 else HC-2
                        back=SP.ZD(BM,H,abs(L),data,i,Hsx) if abs(L)!=0 else SP.ZD(BM,H,abs(L+2),data,i,Hsx)
                        log(back)
                        Hsx=back[0]
                        i=back[2]
                        zhq=zh
                        Hq=HS        
                    data.append([f'{SP.mileageToStr(zh)}\n{pianju}','','',round(Hsx-round(HS+pc/1000,3),3),'',round(HS+pc/1000,3),HS,pc])
        
            log(Hq,BM,zhq,i,dest)
            #dest=random.uniform(0.6,1.5) if Hq-BM[2]>0 else random.uniform(2.6,4.5)
            
            H=Hsx-BM[2]-dest
            L=BM[0]-zhq
            back=SP.ZD(BM,H,abs(L),data,i,Hsx) if abs(L)!=0 else SP.ZD(BM,H,abs(L+2),data,i,Hsx)
            Hsx=back[0]
            if data!=[]:
                Ep=random.randint(-1*i,1*i)
                data.append([f'{BM[1]}','','',round(Hsx-BM[2],3),'',round(BM[2]+Ep/1000,3),BM[2],Ep])            
                DATA.append([data,[round(BM[2]+Ep/1000,3),BM[2],Ep]])
            kwargs['data']=DATA    
            SP.writeGC(app,**kwargs)
        except Exception as err:
            log(err)
        else:
            f=open('高程信息.txt','a+',encoding='UTF-8')
            f.write(f"{time.strftime('%Y-%m-%d %H:%M:%S', time.localtime())}  :  {kwargs['zhuanghao']} 第{kwargs['NO']}层 回填 成功\n")
            f.close()            
        finally:
            app.quit()
        pass

    def PMWZbat(self,DD,**kwargs):
        global ce
        
        try:
            app=xw.App(visible=False, add_book=False)

            path=f"{kwargs['path']}\\{kwargs['gongcheng']}\\{kwargs['gongxu']}"
            
            if os.path.isdir(path):
                pass
            else:
                os.makedirs(path)
            log(path)
            kwargs['path']=path
            HEAD=None
            DATA=[]
            data=[]
            No=1
            condtions=['回填','路床','路提','基层','底基层','沥青']
            kwargs['zhuanghao']=f"{kwargs['buwei']} {kwargs['side'] if 'side' in kwargs else ''}" if all(con not in kwargs['gongxu'] for con in condtions) else f"{kwargs['buwei']} "

            for row in DD:
                
                zh=round(float(row[0]),3)
                pianju=PJFM(row[1])
                ce.SET(zh,kwargs['rw'])
                point=ce.Point(row[1])
                X=round(point[0],4)
                Y=round(point[1],4)
                pc=int(round(((kwargs['pw']**2)/2)**0.5))
                px=random.randint(-pc,pc)
                py=random.randint(-pc,pc)
                ps=round((px**2+py**2)**0.5)
                args=()
                if HEAD!=SP.HEADER(OPTION.getCtr(zh),*args):

                    if data!=[]:
                        log('data',data)
                        #kwargs['data']=data
                        #kwargs['header']=HEAD
                        #kwargs['No']=No
                        DATA.append([data,HEAD])
                        data=[]
                    
                    data.append([f'{SP.mileageToStr(zh)} {pianju}',None,X,None,Y,None,round(X+px/1000,4),None,round(Y+py/1000,4),None,None,px,None,py,None,ps,None])
                else:
                    
                    data.append([f'{SP.mileageToStr(zh)} {pianju}',None,X,None,Y,None,round(X+px/1000,4),None,round(Y+py/1000,4),None,None,px,None,py,None,ps,None])
                HEAD=SP.HEADER(OPTION.getCtr(zh),*args)

            if data!=[]:
                log('data',data)
                DATA.append([data,HEAD])
            kwargs['data']=DATA
            SP.writePM(app,**kwargs)
        except Exception as err:
            log(err)
        else:
            f=open('平面位置信息.txt','a+',encoding='UTF-8')
            f.write(f"{time.strftime('%Y-%m-%d %H:%M:%S', time.localtime())}  :  {kwargs['zhuanghao']} 平面位置 成功\n")
            f.close()
        finally:
       
            app.quit()
            pass

    def FXbat(self,DD,**kwargs):
        global ce

        try:
            app=xw.App(visible=False, add_book=False)
            
            path=f"{kwargs['path']}\\{kwargs['gongcheng']}\\{kwargs['gongxu']}"

            if os.path.isdir(path):
                pass
            else:
                os.makedirs(path)
            log(path)
            kwargs['path']=path
            DATA=[]
            HEAD=None
            data=[]
            No=1
            condtions=['回填','路床','路提','基层','底基层','沥青']
            kwargs['zhuanghao']=f"{kwargs['buwei']} {kwargs['side'] if 'side' in kwargs else ''}" if all(con not in kwargs['gongxu'] for con in condtions) else f"{kwargs['buwei']} "

            for row in DD:
                

                zh=round(float(row[0]),3)
                pianju=PJFM(row[1])

                ce.SET(zh,kwargs['rw'])
                point=ce.Point(row[1])
                X=round(point[0],4)
                Y=round(point[1],4)

                a,s=SP.HEADER(OPTION.getCtr(zh),*(X,Y))
                pc=round(random.uniform(-0.02,0.02),3)
                s1=round(s+pc,3)
                H=round(float(row[2]),3)
                
                if HEAD!=SP.HEADER(OPTION.getCtr(zh)):

                    if data!=[]:
                        log('data',data)
                        DATA.append([data,HEAD])

                        data=[]
                    
                    data.append([f'{SP.mileageToStr(zh)} {pianju}',None,X,None,Y,None,a,None,None,s,None,s1,None,abs(pc),None,H,None])
                else:
                    
                    data.append([f'{SP.mileageToStr(zh)} {pianju}',None,X,None,Y,None,a,None,None,s,None,s1,None,abs(pc),None,H,None])
                HEAD=SP.HEADER(OPTION.getCtr(zh))

            if data!=[]:
                log('data',data)
                DATA.append([data,HEAD])
            kwargs['data']=DATA
            SP.writeFX(app,**kwargs)
        except Exception as e:
            log(e)
        else:
            f=open('放线记录信息.txt','a+',encoding='UTF-8')
            f.write(f"{time.strftime('%Y-%m-%d %H:%M:%S', time.localtime())}  :  {kwargs['zhuanghao']} 放线记录 成功\n")
            f.close()
        finally:
            app.quit()

    def LMbat(self,DD,**kwargs):
        global ce
        try:

            app=xw.App(visible=False, add_book=False)            
            path=f"{kwargs['path']}\\{kwargs['gongcheng']}\\{kwargs['gongxu']}"
            if os.path.isdir(path):
                pass
            else:
                os.makedirs(path)
            kwargs['path']=path
            condtions=['回填','路床','路提','基层','底基层','沥青']
            kwargs['zhuanghao']=f"{kwargs['buwei']} {kwargs['side'] if 'side' in kwargs else ''}" if all(con not in kwargs['gongxu'] for con in condtions) else f"{kwargs['buwei']} "
            kwargs['gongxu']=f"{kwargs['gongxu']} 第{kwargs['NO']}层 共{kwargs['total']}层"
            DATA=[]
            data=[]
            BM=None
            Hsx=None
            dest=None
            Hq=None
            zhq=None
            No=1
            i=0        
            for row in DD:
                
                zh=round(float(row[0]),3)
                pianju=PJFM(row[1])
                HS=round(float(row[2]),3)
                pc=random.randint(kwargs['pc_s'],kwargs['pc_b'])
                if BM!=OPTION.getBM(zh):
                    if data !=[]:
                        H=Hsx-BM[2]-dest
                        L=BM[0]-zhq
                        back=SP.LMZD(BM,H,abs(L),data,i,Hsx) if abs(L)!=0 else SP.LMZD(BM,H,abs(L+2),data,i,Hsx)
                        Hsx=back[0]
                        Ep=random.randint(-1*i,1*i)
                        data.append([f'{BM[1]}','','','',round(Hsx-BM[2],3),'',round(BM[2]+Ep/1000,3),BM[2],Ep,''])
                        
                        DATA.append([data,[round(BM[2]+Ep/1000,3),BM[2],Ep]])
                        data=[]
                    BM=OPTION.getBM(zh) 
                    H=BM[2]-HS
                    L=abs(zh-BM[0])
                    Hsx=None
                    i=0
                    back=SP.LMZD(BM,H,abs(L),data,i,Hsx)
                    Hsx=back[0]
                    i=back[2]
                    dest=back[1]
                    zhq=zh
                    Hq=HS 
                    data.append([f'{SP.mileageToStr(zh)}',f'{pianju}','','',round(Hsx-round(HS+pc/1000,3),3),'',round(HS+pc/1000,3),HS,pc,f"{kwargs['pc_s']},{kwargs['pc_b']}"])
                else:
                    HC=Hsx-HS            
                    L=zh-zhq
                    if HC<0.6 or HC>4.8 or L>30:
                        H=Hsx-HS-dest if HC> 4.8 else HC-2
                        back=SP.LMZD(BM,H,abs(L),data,i,Hsx) if abs(L)!=0 else SP.LMZD(BM,H,abs(L+2),data,i,Hsx)
                        log(back)
                        Hsx=back[0]
                        i=back[2]
                        zhq=zh
                        Hq=HS        
                    data.append([f'{SP.mileageToStr(zh)}',f'{pianju}','','',round(Hsx-round(HS+pc/1000,3),3),'',round(HS+pc/1000,3),HS,pc,f"{kwargs['pc_s']},{kwargs['pc_b']}"])
        
            log(Hq,BM,zhq,i,dest)
            #dest=random.uniform(0.6,1.5) if Hq-BM[2]>0 else random.uniform(2.6,4.5)
            
            H=Hsx-BM[2]-dest
            L=BM[0]-zhq
            back=SP.LMZD(BM,H,abs(L),data,i,Hsx) if abs(L)!=0 else SP.LMZD(BM,H,abs(L+2),data,i,Hsx)
            Hsx=back[0]
            if data!=[]:
                Ep=random.randint(-1*i,1*i)
                data.append([f'{BM[1]}','','','',round(Hsx-BM[2],3),'',round(BM[2]+Ep/1000,3),BM[2],Ep,''])
                for i in data:
                    print(len(i),i)
                DATA.append([data,[round(BM[2]+Ep/1000,3),BM[2],Ep]])
            kwargs['data']=DATA    
            SP.writeGCJS(app,**kwargs)
        except Exception as err:
            log(err)
        else:
            f=open('高程信息.txt','a+',encoding='UTF-8')
            f.write(f"{time.strftime('%Y-%m-%d %H:%M:%S', time.localtime())}  :  {kwargs['zhuanghao']} 第{kwargs['NO']}层 回填 成功\n")
            f.close()            
        finally:
            app.quit()
        pass

    ################################################################


APP()





    
